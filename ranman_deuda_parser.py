"""
ranman_deuda_parser.py — RANMAN "Cuadro de Riesgos" (debt with cost) parser.

A faithful port of Valoran's `actualizar_cuadro_riesgos.py` (Sep-2026 version,
kept in OneDrive under Archivos Ranman/_Tablero de riesgo). RANMAN publishes a
debt snapshot ("corte") every week or two as `Cuadro de riesgos Ranman <dd mmm
aa>.xlsm`; this module reads the `deuda con costo` sheet of one such file and
returns the credit-level detail plus the per-corte summary record, using the
SAME classification rules so Maquina's series is comparable with the tablero's:

  • Puentes      — debt secured by the project's own land (BBVA / BANREGIO /
                   BanCrea bridge loans, Azhala land credit, the Cúspide
                   infrastructure share of the BanBajío line). Repaid from the
                   homes they finance.
  • No-Puentes   — corporate lines (CCC 55 mdp, BBVA simples 99.8 / 100 / 28,
                   BanCrea línea 48, the corporate share of BanBajío, the VRM
                   credit). No own source of repayment — the block that matters.

  The tablero splits the book three ways (`tipo3`): Puentes · Créditos Simples
  (amortised principal cannot be redrawn: the BBVA simples, BanBajío's
  corporate share, the VRM credit) · Líneas Revolventes (everything else).

  • Crédito VRM  — debt with the shareholder, NOT in RANMAN's workbooks. It is
                   kept apart in `Credito VRM.xlsx` (a list of dated movements;
                   the balance at a corte is the sum of movements up to that
                   date) and injected into every corte AFTER the corte has
                   reconciled against its own Total, which is bank debt only.

Amounts are in THOUSANDS of MXN, as in the workbook (÷1000 → mdp).

Blocks per bank are located by the text `Fecha de Firma` in column B — never by
row number. A credit is any row with a name and a numeric saldo deudor (some
rows lack a signing date). Each file is reconciled against the Total it declares
on the same sheet; a mismatch is reported to the caller.

Also parses the tablero's `serie.json` / `corte_actual.json` (bulk history
import) and the exported tablero .html (its embedded series).
"""
from __future__ import annotations

import datetime as dt
import io
import json
import re
import unicodedata
from collections import defaultdict

import openpyxl

MESES = {'ene': 1, 'feb': 2, 'mzo': 3, 'mar': 3, 'abr': 4, 'may': 5, 'jun': 6,
         'jul': 7, 'ago': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dic': 12}

# Créditos simples: amortised principal can NOT be redrawn. BanBajío's
# infrastructure line amortises to 2033, so it counts as simple.
SIMPLES = {'Simple 99.8 mdp', 'Simple 100 mdp', 'Simple 28 mdp', 'Infraestructura', 'Crédito VRM'}
# The three BBVA simples alone (the earlier two-way "Amortizable" definition;
# still used to read records written before the three-way split existed).
AMORTIZABLES = {'Simple 99.8 mdp', 'Simple 100 mdp', 'Simple 28 mdp'}
TIPOS3 = ('Puentes', 'Créditos Simples', 'Líneas Revolventes')

BANCO = {'Hipotecaria Nacional': 'BBVA', 'BBVA CCC 55MDP': 'BBVA',
         'BBVA SIMPLE 100mdp (99.780 MDP)': 'BBVA', 'BBVA SIMPLE 100 MDP': 'BBVA',
         'BBVA SIMPLE 28MDP': 'BBVA', 'BanBajio': 'BanBajío', 'BANREGIO': 'BANREGIO',
         'BanCrea': 'BanCrea', 'ION': 'ION', 'SANTANDER': 'SANTANDER'}

# Authorised size of the revolving No-Puentes lines (mdp), read off the line names.
REVOLVING_AUTHORIZED_MDP = {'CCC 55 mdp': 55.0, 'Línea 48 mdp': 48.0}

# Lines the script leaves under their raw group name; the tablero names them.
LINEA_ALIAS = {'BANREGIO': 'Línea BANREGIO', 'SANTANDER': 'Línea SANTANDER',
               'BanCrea': 'Línea 48 mdp', 'ION': 'Línea ION'}
# (tipo, banco, authorised amount in mdp — None when not disclosed), as the tablero has them.
LINEAS_META = {
    'CCC 55 mdp':      ('Líneas Revolventes', 'BBVA',      55.0),
    'Crédito VRM':     ('Créditos Simples',   'VRM',       None),   # shareholder debt: no line, no maturity
    'Infraestructura': ('Créditos Simples',   'BanBajío',  15.0),   # corporate share; Cúspide's share is a Puente
    'Línea 48 mdp':    ('Líneas Revolventes', 'BanCrea',   48.0),
    'Línea BANREGIO':  ('Líneas Revolventes', 'BANREGIO',  None),
    'Línea SANTANDER': ('Líneas Revolventes', 'SANTANDER', None),
    'Simple 100 mdp':  ('Créditos Simples',   'BBVA',     100.0),
    'Simple 28 mdp':   ('Créditos Simples',   'BBVA',      28.0),
    'Simple 99.8 mdp': ('Créditos Simples',   'BBVA',     100.0),
}


def _ascii(s: str) -> str:
    return unicodedata.normalize('NFKD', s.lower()).encode('ascii', 'ignore').decode()


def fecha_de_nombre(nombre: str):
    """'Cuadro de riesgos Ranman 31 ago 26.xlsm' -> '2026-08-31' (ISO) or None."""
    s = nombre.lower().replace('.xlsm', '').replace('.xlsx', '').replace('_modificado', '').replace(' vf', '')
    m = re.search(r'(\d{1,2})\s*([a-z]{3})\s*(\d{2})\s*$', s)
    if not m:
        return None
    mes = MESES.get(m.group(2))
    if not mes:
        return None
    try:
        return dt.date(2000 + int(m.group(3)), mes, int(m.group(1))).isoformat()
    except ValueError:
        return None


def desarrollo(n: str) -> str:
    s = _ascii(n)
    for k, v in [('puerta natura', 'Puerta Natura'), ('puerta de piedra', 'Corporativo'),
                 ('cantabria', 'Cantabria'), ('trive', 'Trive'), ('cartuja', 'Cartuja'),
                 ('cuspide', 'Cúspide'), ('alella', 'Alella'), ('azhala', 'Azhala'),
                 ('carmel', 'Carmel'), ('miranda', 'Celaya'), ('celaya', 'Celaya'),
                 ('gran natura', 'Corporativo'), ('mompani', 'Corporativo'),
                 ('aurea', 'Corporativo'), ('veredas', 'Corporativo'),
                 ('punta norte', 'Corporativo'), ('ximonco', 'Corporativo'),
                 ('inmuebles', 'Corporativo'), ('10k25', 'Corporativo')]:
        if k in s:
            return v
    return 'Corporativo'


def clasifica(grupo: str, nombre: str, desa: str, garantia: str = ''):
    """Same rule as lámina 21: Puentes = debt secured by the project's own land;
    No-Puentes = corporate lines. Returns (línea, categoría)."""
    g = grupo.upper()
    s = _ascii(nombre)
    if 'tierra' in s:
        return 'Crédito tierra', 'Puentes'
    if grupo == 'Hipotecaria Nacional':
        return 'Crédito puente', 'Puentes'
    if re.match(r'^(pre)?p[ru]{1,2}ente', s) or 'puente' in s.split('.-')[0]:
        return 'Crédito puente', 'Puentes'
    if 'CCC' in g:
        return 'CCC 55 mdp', 'No-Puentes'
    if '99.780' in g or '99.78' in g:
        return 'Simple 99.8 mdp', 'No-Puentes'
    if g == 'BBVA SIMPLE 100 MDP':
        return 'Simple 100 mdp', 'No-Puentes'
    if '28MDP' in g:
        return 'Simple 28 mdp', 'No-Puentes'
    if g == 'BANBAJIO':
        # Shared line: the Cúspide infrastructure share is project debt; the rest corporate.
        return ('Infraestructura', 'Puentes') if desa == 'Cúspide' else ('Infraestructura', 'No-Puentes')
    if g == 'BANCREA':
        # Until Nov-2025 BanCrea's credits came without the "Puente.-" /
        # "linea48mdp.-" prefix. The collateral always tells them apart: the
        # bridge loan is secured with the project's own land, the 48 mdp line
        # with Puerta Natura's commercial lots.
        if 'del proyecto' in (garantia or '').lower():
            return 'Crédito puente', 'Puentes'
        return 'Línea 48 mdp', 'No-Puentes'
    if g == 'BANREGIO' and 'linea' in s:
        return 'Línea BANREGIO', 'No-Puentes'
    return grupo, 'No-Puentes'


def tipo3_de(linea: str, cat: str) -> str:
    """The tablero's three-way class."""
    if cat == 'Puentes':
        return 'Puentes'
    if linea in SIMPLES:
        return 'Créditos Simples'
    return 'Líneas Revolventes'


def tipo3(c: dict) -> str:
    return tipo3_de(c['linea'], c.get('cat') or ('Puentes' if c['linea'] in ('Crédito puente', 'Crédito tierra') else 'No-Puentes'))


def parse_corte(file_bytes: bytes, filename: str = ''):
    """One corte workbook -> (fecha_iso, creditos, declarado_total). Bank debt
    only — the VRM credit is added by the caller once the file reconciles.

    Raises ValueError when the file isn't a Cuadro de Riesgos."""
    fecha = fecha_de_nombre(filename or '')
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        if 'deuda con costo' not in wb.sheetnames:
            raise ValueError("'%s' has no 'deuda con costo' sheet — is it a Cuadro de Riesgos?" % filename)
        filas = list(wb['deuda con costo'].iter_rows(min_row=1, max_row=200, max_col=17, values_only=True))
    finally:
        wb.close()

    declarado = None
    for r in filas:
        if r[0] and str(r[0]).strip().lower() == 'total' and isinstance(r[1], (int, float)):
            declarado = float(r[1])
            break

    creds, grupo = [], None
    for r in filas:
        if r[1] == 'Fecha de Firma' and r[0]:
            grupo = str(r[0]).strip()
            continue
        if not grupo or not r[0]:
            continue
        nombre = str(r[0]).strip()
        if nombre.upper().startswith('TOTAL'):
            grupo = None
            continue
        if not isinstance(r[7], (int, float)):        # no saldo deudor -> not a credit
            continue
        desa = desarrollo(nombre)
        garantia = str(r[16]) if len(r) > 16 and r[16] else ''
        linea, cat = clasifica(grupo, nombre, desa, garantia)
        creds.append(dict(
            banco=BANCO.get(grupo, grupo), credito=nombre, linea=linea, cat=cat,
            desarrollo=desa, saldo=float(r[7]),
            vence=r[2].date().isoformat() if isinstance(r[2], dt.datetime) else (r[2].isoformat() if isinstance(r[2], dt.date) else None),
            porEjercer=float(r[8]) if isinstance(r[8], (int, float)) else 0.0,
            sobretasa=float(r[14]) if isinstance(r[14], (int, float)) else None,
            avance=float(r[13]) if isinstance(r[13], (int, float)) else None))
    if not creds:
        raise ValueError("'%s': no credit rows found under any 'Fecha de Firma' block." % filename)
    return fecha, creds, declarado


# ── the VRM credit (Credito VRM.xlsx) ───────────────────────────────────────
MESES_ES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
            'julio': 7, 'agosto': 8, 'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12}


def fecha_es(v):
    """«24 Mayo 2024» -> «2024-05-24». Also accepts a real Excel date."""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    m = re.match(r'\s*(\d{1,2})\s+([A-Za-zÁÉÍÓÚáéíóúñÑ]+)\s+(\d{4})\s*$', str(v or ''))
    if not m:
        return None
    mes = MESES_ES.get(m.group(2).lower())
    if not mes:
        return None
    try:
        return dt.date(int(m.group(3)), mes, int(m.group(1))).isoformat()
    except ValueError:
        return None


def parse_vrm(file_bytes: bytes) -> list:
    """`Credito VRM.xlsx` -> sorted [[fecha_iso, monto_miles], ...]. Each row of
    the file is one movement (column B date, column C amount in pesos); the
    balance at a date is the sum of the movements up to it."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        flujos = []
        for fila in ws.iter_rows(min_col=2, max_col=3, values_only=True):
            cuando, monto = fila[0], fila[1]
            if not isinstance(monto, (int, float)):
                continue
            f = fecha_es(cuando)
            if f:
                flujos.append([f, round(monto / 1000.0, 6)])
    finally:
        wb.close()
    if not flujos:
        raise ValueError('Credito VRM.xlsx: no dated movements found (expected a date in column B and an amount in column C).')
    return sorted(flujos)


def credito_vrm(flujos: list, fecha: str) -> list:
    """The VRM credit at a corte, shaped like any other credit (thousands)."""
    saldo = sum(m for f, m in (flujos or []) if f <= fecha)
    if abs(saldo) < 0.005:
        return []
    return [dict(banco='VRM', credito='Crédito VRM', linea='Crédito VRM',
                 cat='No-Puentes', desarrollo='Corporativo', saldo=round(saldo, 2),
                 vence=None,            # no agreed maturity
                 porEjercer=0, sobretasa=0.03, avance=None, noBancario=True)]


# ── per-corte record ─────────────────────────────────────────────────────────
def _agrupa(creds, campo):
    d = defaultdict(float)
    for c in creds:
        d[c[campo]] += c['saldo']
    return {k: round(v, 2) for k, v in d.items() if abs(v) > 0.005}


def serie_record(fecha: str, archivo: str, creds: list) -> dict:
    """The per-corte summary — identical shape to the tablero's serie.json.
    Tags every credit with its tipo3 (as corte_actual.json carries it)."""
    for c in creds:
        c['tipo3'] = tipo3(c)
    total = sum(c['saldo'] for c in creds)
    npc = [c for c in creds if c['cat'] == 'No-Puentes']
    corte12 = dt.date.fromisoformat(fecha) + dt.timedelta(days=365)
    consob = [c for c in creds if c['sobretasa'] and c['saldo']]
    base_sob = sum(c['saldo'] for c in consob)
    no_puente = [c for c in creds if c['tipo3'] != 'Puentes']
    return dict(
        fecha=fecha, archivo=archivo, total=round(total, 2), n=len(creds),
        cat=_agrupa(creds, 'cat'), banco=_agrupa(creds, 'banco'),
        desarrollo=_agrupa(creds, 'desarrollo'), linea=_agrupa(creds, 'linea'),
        np={'Amortizable': round(sum(c['saldo'] for c in npc if c['linea'] in SIMPLES), 2),
            'Revolvente': round(sum(c['saldo'] for c in npc if c['linea'] not in SIMPLES), 2)},
        tipo3=_agrupa(creds, 'tipo3'),
        # per-line series, non-puente part only: follows how each simple and
        # each revolving line is drawn or repaid
        lineasNP=_agrupa(no_puente, 'linea'),
        lineasTipo={c['linea']: c['tipo3'] for c in no_puente},
        porEjercer=round(sum(c['porEjercer'] for c in creds), 2),
        sobretasaPond=round(sum(c['sobretasa'] * c['saldo'] for c in consob) / base_sob, 6) if base_sob else None,
        vence12=round(sum(c['saldo'] for c in creds
                          if c['vence'] and dt.date.fromisoformat(c['vence']) <= corte12), 2))


def reconcile(creds: list, declarado) -> str | None:
    """None when the parsed sum matches the sheet's own Total (±1 thousand).
    Bank credits only — call before adding the VRM credit."""
    total = sum(c['saldo'] for c in creds if not c.get('noBancario'))
    if declarado is not None and abs(total - declarado) > 1:
        return 'sum %.1f vs declared Total %.1f mdp' % (total / 1000, declarado / 1000)
    return None


# ── reading stored records (any vintage) ─────────────────────────────────────
def _infra_corp(rec: dict) -> float:
    """Corporate share of BanBajío's shared line, from a record without tipo3:
    whatever part of Puentes isn't a crédito puente/tierra is Cúspide's share."""
    lin = rec.get('linea') or {}
    cat = rec.get('cat') or {}
    infra = float(lin.get('Infraestructura') or 0)
    resto = float(cat.get('Puentes') or 0) - float(lin.get('Crédito puente') or 0) - float(lin.get('Crédito tierra') or 0)
    return infra - min(infra, max(0.0, resto))


def tres_vias_de_record(rec: dict) -> dict:
    """Three-way split of a stored corte (thousands): the record's own tipo3
    when it has one, else derived from the two-way record."""
    t3 = rec.get('tipo3')
    if isinstance(t3, dict) and t3:
        return {t: round(float(t3.get(t) or 0), 2) for t in TIPOS3}
    lin = rec.get('linea') or {}
    pu = float((rec.get('cat') or {}).get('Puentes') or 0)
    si = sum(float(lin.get(s) or 0) for s in AMORTIZABLES) + _infra_corp(rec) + float(lin.get('Crédito VRM') or 0)
    re_ = float(rec.get('total') or 0) - pu - si
    return {'Puentes': round(pu, 2), 'Créditos Simples': round(si, 2), 'Líneas Revolventes': round(max(0.0, re_), 2)}


def lineas_de_record(rec: dict) -> dict:
    """Per-line non-puente saldos of a stored corte (thousands), under the
    tablero's names."""
    out = {}
    lnp = rec.get('lineasNP')
    if isinstance(lnp, dict):
        for k, v in lnp.items():
            name = LINEA_ALIAS.get(k, k)
            out[name] = out.get(name, 0.0) + float(v or 0)
        return {k: round(v, 2) for k, v in out.items()}
    for k, v in (rec.get('linea') or {}).items():
        if k in ('Crédito puente', 'Crédito tierra'):
            continue
        val = _infra_corp(rec) if k == 'Infraestructura' else float(v or 0)
        name = LINEA_ALIAS.get(k, k)
        out[name] = out.get(name, 0.0) + val
    return {k: round(v, 2) for k, v in out.items()}


def lineas_tipo_de_record(rec: dict) -> dict:
    """{line name: tipo3} as the record carries it (empty for older records)."""
    lt = rec.get('lineasTipo')
    if not isinstance(lt, dict):
        return {}
    return {LINEA_ALIAS.get(k, k): v for k, v in lt.items()}


# ── tablero exports ──────────────────────────────────────────────────────────
def _js_const(txt: str, name: str):
    """The JSON literal assigned to `const <name> = ...;` on one line of the export."""
    key = 'const %s = ' % name
    i = txt.find(key)
    if i < 0:
        return None
    j = txt.find('\n', i)
    return json.loads(txt[i + len(key):j].strip().rstrip(';'))


def parse_tablero_html(file_bytes: bytes) -> dict:
    """The exported tablero ('Cuadro de Riesgo RANMAN.html') embeds its data:
    `const SD = {cortes, lineas}` in mdp and `const DATA = {asOf, creditos}` in
    thousands. Returns dict(cortes={fecha: {tipo3, lineasNP}} in THOUSANDS,
    actual=DATA or None, lineas={name: (tipo, banco, aut)})."""
    txt = file_bytes.decode('utf-8', errors='replace')
    sd = _js_const(txt, 'SD')
    if not isinstance(sd, dict) or not isinstance(sd.get('cortes'), list) or not sd['cortes']:
        raise ValueError('The .html is not a tablero export with an embedded SD (cortes) block.')
    lineas_sd = sd.get('lineas') if isinstance(sd.get('lineas'), dict) else {}
    cortes = {}
    for i, x in enumerate(sd['cortes']):
        f = x.get('f')
        dt.date.fromisoformat(f)                     # validates
        t3 = {t: round(float((x.get('c') or {}).get(t) or 0) * 1000.0, 2) for t in TIPOS3}
        lin, ltipo = {}, {}
        for name, d in lineas_sd.items():
            vals = d.get('v') or []
            v = float(vals[i]) if i < len(vals) and vals[i] is not None else 0.0
            if v > 0.0005:
                lin[name] = round(v * 1000.0, 2)
                ltipo[name] = d.get('tipo') or tipo3_de(name, 'No-Puentes')
        cortes[f] = {'tipo3': t3, 'lineasNP': lin, 'lineasTipo': ltipo}
    actual = _js_const(txt, 'DATA')
    if isinstance(actual, dict) and actual.get('asOf') and isinstance(actual.get('creditos'), list):
        for c in actual['creditos']:                 # the page's shape back to the script's
            if c.get('noBanc') and not c.get('noBancario'):
                c['noBancario'] = True
            c.setdefault('cat', 'Puentes' if c.get('linea') in ('Crédito puente', 'Crédito tierra')
                         or (c.get('linea') == 'Infraestructura' and c.get('desarrollo') == 'Cúspide') else 'No-Puentes')
    else:
        actual = None
    meta = {name: (d.get('tipo'), d.get('banco'), d.get('aut')) for name, d in lineas_sd.items()}
    return dict(cortes=cortes, actual=actual, lineas=meta)


def parse_serie_json(file_bytes: bytes) -> list:
    data = json.loads(file_bytes.decode('utf-8'))
    if not isinstance(data, list) or not data or not all(isinstance(r, dict) and 'fecha' in r and 'total' in r for r in data):
        raise ValueError('serie.json should be a list of corte records with fecha and total.')
    out = []
    for r in data:
        dt.date.fromisoformat(r['fecha'])          # validates
        out.append(r)
    return out


def parse_corte_actual_json(file_bytes: bytes) -> dict:
    data = json.loads(file_bytes.decode('utf-8'))
    if not isinstance(data, dict) or 'asOf' not in data or not isinstance(data.get('creditos'), list):
        raise ValueError('corte_actual.json should carry asOf and a creditos list.')
    dt.date.fromisoformat(data['asOf'])
    return data
