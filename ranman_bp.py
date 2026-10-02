"""RANMAN's business plans «NN BP <PLAZA> V02.xlsx» -> unit economics by etapa.

Ranman keeps one financial model per plaza (SLP also has Kima and Trivé on
their own, Ximonco is Cantabria) and drops the current version into the month
folder, «5 BPs/<NN BP plaza>/». Every model has the same «Supuestos» sheet: one
block per etapa, its P&L three ways side by side —

    Proforma firmado   what was signed when the etapa was approved (the UW)
    Cierre previsto    where Ranman expects it to close today
    Resultados FLUJOS  what the model's monthly cash flow actually adds up to

— with the etapa's start and end dates, its number of units (lotes,
departamentos), the price and margin of each prototype, and how it is funded.
Row 2 says up to which month the model carries actuals («INFORMACIÓN
HISTÓRICA»): that month is the cut. «Saldos Finales» has every etapa's
month-end cash balance.

Nothing is read by fixed rows: a block starts at an etapa name in column B with
«Proforma firmado» right below it, its three amount columns are the «Monto»
headers, and the lines are found by label (the models have drifted — the
Jalisco blocks say «Gasto Admin» twice, the San Luis model keeps its unit count
ten columns further right).

Checked before anything is stored: each etapa's P&L adds up (sales less costs
= gross margin, less expenses = EBIT, less financial cost = UAIR) in the
expected-close and cash-flow columns; the signed proforma is typed in by hand,
so a miss there is only reported. The «Resumen» sheet is not used as a check:
each plaza builds its own (Celaya's sums Miranda only, Aguascalientes splits
homes from lots). Amounts come in thousands of pesos and are stored in mdp.
"""
import calendar
import datetime as dt
import io
import re

import openpyxl

from ranman_package import _norm, _num, _hoja, clave_desarrollo

# The P&L lines, in the order of the Edo. Resultados (ranman_package.CONCEPTOS)
# plus the mezzanine cost the BPs carry on its own line.
CONCEPTOS = [('ventas', 'Ventas'), ('terreno', 'Terreno'), ('permisos', 'Permisos y Licencias'),
             ('estudios', 'Estudios y Proyectos'), ('urbanizacion', 'Urbanización'),
             ('infraestructura', 'Infraestructura'), ('edificacion', 'Edificación'),
             ('margen_bruto', 'Margen Bruto'), ('admin_plaza', 'Admin (Plaza)'),
             ('admin_corp', 'Admin (Corporativo)'), ('comercial', 'Comercial'), ('comisiones', 'Comisiones'),
             ('mantenimiento', 'Mantenimiento'), ('ebit', 'EBIT'), ('fin_operativos', 'Financieros Operativos'), ('fin_lp', 'Financieros LP'),
             ('fin_mezz', 'Financieros Mezzanine'), ('uair', 'UAIR')]
COSTOS = ('terreno', 'permisos', 'estudios', 'urbanizacion', 'infraestructura', 'edificacion')
GASTOS = ('admin_plaza', 'admin_corp', 'comercial', 'comisiones', 'mantenimiento')
FINANCIEROS = ('fin_operativos', 'fin_lp', 'fin_mezz')
# Line labels, compared with spaces and punctuation taken out: the models spell
# them a few ways («Gasto AdminCorp.», «Financieros Largo Plazo», the Cantabria
# model's «Financieros Inversionista» for the mezzanine), and the 2025 models
# split or merge some lines (a single «Financieros», «Infraestructura Interna /
# Externa», «Infraestructura compartida» on top of «Infraestructura»). Two
# different labels that land on the same line are added up.
_ETIQUETA = {'ventastotales': 'ventas', 'ventas': 'ventas', 'terreno': 'terreno',
             'permisosylicencias': 'permisos', 'estudiosyproyectos': 'estudios',
             'urbanizacion': 'urbanizacion', 'infraestructura': 'infraestructura', 'edificacion': 'edificacion',
             'margenbruto': 'margen_bruto', 'gastoadminplaza': 'admin_plaza', 'gastoadmplaza': 'admin_plaza',
             'gastoadmincorp': 'admin_corp', 'gastoadmcorp': 'admin_corp', 'gastoadmincorporativo': 'admin_corp',
             'gastocomercial': 'comercial', 'comisionesventa': 'comisiones', 'comisiones': 'comisiones',
             'ebit': 'ebit', 'financierosoperativos': 'fin_operativos', 'financieroslp': 'fin_lp',
             'financieroslargoplazo': 'fin_lp', 'financierosmezzanine': 'fin_mezz',
             'financierosinversionista': 'fin_mezz', 'financieros': 'fin_operativos',
             'infraestructurainterna': 'urbanizacion', 'infraestructuraexterna': 'infraestructura',
             'infraestructuracompartida': 'infraestructura', 'mantenimiento': 'mantenimiento', 'uair': 'uair'}


def _concepto(v):
    return _ETIQUETA.get(re.sub(r'[^a-z]', '', _norm(v))) if isinstance(v, str) else None


ESCENARIOS = [('proforma', 'Proforma firmado'), ('cierre', 'Cierre previsto'), ('flujos', 'Resultados Flujos')]
_FONDEO = {'necesidad de equity': 'equity', 'credito puente': 'puente', 'credito corporativo': 'corporativo',
           'credito mezzanine': 'mezzanine'}
_PROTO = {'prototipo': 'n', 'mezcla': 'mezcla', 'm2 terreno': 'm2t', 'm2 construccion': 'm2c',
          'precio promedio': 'precio', 'unidades': 'u', 'uair': 'uair', 'vendidas': 'vendidas'}


class SinSupuestos(ValueError):
    """A BP-named workbook without the per-etapa sheet (the 2024 consolidated
    «00 BP RANMAN», a land study): recognised, but nothing to read."""


def _mdp(v):
    n = _num(v)
    return round(n / 1000.0, 3) if n is not None else None


def _mes(v):
    return '%d-%02d' % (v.year, v.month) if isinstance(v, (dt.datetime, dt.date)) else None


def _fin_de_mes(ym):
    y, m = (int(x) for x in ym.split('-'))
    return dt.date(y, m, calendar.monthrange(y, m)[1]).isoformat()


def slug_de(filename):
    """«00 BP SAN LUIS POTOSÍ V02.xlsx» -> 'san_luis_potosi', «BP Punta Norte V4 - VT» -> 'punta_norte',
    «BP MOMPANI VF» -> 'mompani': the model's name without its plaza number and
    version, so V02, V03 and the «versión final» are one model."""
    n = _norm(re.split(r'[\\/]', filename or '')[-1])
    n = re.sub(r'\.xls[xm]$', '', n)
    n = re.sub(r'^\d{1,2} ', '', n)
    n = re.sub(r'^bp ', '', n)
    n = re.sub(r'\s+-\s+.*$', '', n)                  # « - VT», « - copia»
    n = re.sub(r'\s+(v\d+(\.\d+)?|vf)\b.*$', '', n)   # « V02», « V4», « VF»
    return re.sub(r'[^a-z0-9]+', '_', n).strip('_') or 'bp'


def nombre_de(filename):
    """«00 BP SAN LUIS POTOSÍ V02.xlsx» -> 'SAN LUIS POTOSÍ V02' (shown on the page)."""
    n = re.split(r'[\\/]', filename or '')[-1]
    n = re.sub(r'\.xls[xm]$', '', n, flags=re.I)
    return re.sub(r'^(\d{1,2} )?bp ', '', n, flags=re.I).strip()


def es_estudio(ruta):
    """True for a model in a folder of its own rather than the plaza's
    («5 BPs/02 BP Aguascalientes/Viñedos Elizondo 08-26/BP Viñedos Elizondo V4.xlsx»):
    a project under study. The plaza models sit in «NN BP <plaza>», in «5 BPs»
    itself or in a dated «BPs al dd-mm-aa». Without a path, not a study."""
    partes = [p for p in re.split(r'[\\/]+', ruta or '') if p]
    if len(partes) < 2:
        return False
    padre = _norm(partes[-2])
    return not re.match(r'^(\d{1,2}\.? )?bps?( |$)', padre)


# ── «Supuestos»: one block per etapa ────────────────────────────────────────
def _bloques(filas):
    """[(first row, end row)] of every etapa block: a name in column B with
    «Proforma firmado» in column B one to three rows below."""
    inicios = []
    for i, r in enumerate(filas):
        if len(r) > 1 and isinstance(r[1], str) and r[1].strip() and 'proforma' not in _norm(r[1]):
            if any(len(filas[k]) > 1 and 'proforma' in _norm(filas[k][1])
                   for k in range(i + 1, min(i + 4, len(filas)))):
                inicios.append(i)
    return [(a, inicios[n + 1] if n + 1 < len(inicios) else len(filas)) for n, a in enumerate(inicios)]


def _a_la_derecha(fila, j, cuantos=4, tipo=(int, float)):
    """The first value of the given type in the cells right of column j."""
    for v in fila[j + 1:j + 1 + cuantos]:
        if isinstance(v, tipo) and not isinstance(v, bool):
            return v
    return None


def _lee_bloque(filas, a, b):
    cab = filas[a]
    out = {'n': str(cab[1]).strip(), 'k': clave_desarrollo(cab[1]),
           'estatus': str(cab[4]).strip() if len(cab) > 4 and isinstance(cab[4], str) else None}
    # the three amount columns: the «Monto» headers, left to right
    hm = next((i for i in range(a, min(a + 8, b)) if sum(1 for v in filas[i] if _norm(v) == 'monto') >= 3), None)
    if hm is None:
        raise ValueError("etapa %s: no row with the three «Monto» headers" % out['n'])
    cols = [j for j, v in enumerate(filas[hm]) if _norm(v) == 'monto'][:3]
    esc = {e: {} for e, _ in ESCENARIOS}
    admin_vistos, vistas = 0, set()
    fin_pyl = None
    for i in range(hm + 1, b):
        r = filas[i]
        et = r[1] if len(r) > 1 else None
        compacta = re.sub(r'[^a-z]', '', _norm(et)) if isinstance(et, str) else ''
        if compacta == 'gastoadmin':
            k = ('admin_plaza', 'admin_corp')[min(admin_vistos, 1)]     # Jalisco: plaza first, then corporate
            admin_vistos += 1
        else:
            k = _concepto(et)
            if not k or compacta in vistas:
                continue
            vistas.add(compacta)
        for (e, _), j in zip(ESCENARIOS, cols):
            v = _mdp(r[j]) if j < len(r) else None
            if esc[e].get(k) is not None and v is not None:
                v = round(esc[e][k] + v, 3)
            elif v is None:
                v = esc[e].get(k)
            esc[e][k] = v
        if k == 'uair':
            fin_pyl = i
            break
    if fin_pyl is None:
        raise ValueError("etapa %s: no UAIR line" % out['n'])
    for e in esc:
        for k, _ in CONCEPTOS:
            esc[e].setdefault(k, None)
    out['esc'] = esc

    # the side panel: dates, units, prototypes (anywhere in the block, by label)
    fechas, unidades, protos = {}, None, []
    for i in range(a, b):
        r = filas[i]
        for j, v in enumerate(r):
            t = _norm(v) if isinstance(v, str) else ''
            if not t:
                continue
            if t in ('fecha de inicio', 'fecha de termino') and not fechas.get(t):
                ds = [_mes(x) for x in r[j + 1:j + 6] if isinstance(x, (dt.datetime, dt.date))]
                if ds:
                    fechas[t] = ds[:2]
            elif t.startswith('no. de ') and unidades is None:
                n = _a_la_derecha(r, j)
                if n is not None:
                    unidades = (round(float(n), 2), t[7:].strip())
            elif t == 'prototipo' and not protos:
                hdr = {jj: _PROTO[_norm(x)] for jj, x in enumerate(r) if isinstance(x, str) and _norm(x) in _PROTO}
                for r2 in filas[i + 1:b]:
                    nom = r2[j] if j < len(r2) else None
                    if not isinstance(nom, str) or not nom.strip() or nom.strip().startswith('*'):
                        break
                    p = {}
                    for jj, campo in hdr.items():
                        x = r2[jj] if jj < len(r2) else None
                        p[campo] = x.strip() if campo == 'n' and isinstance(x, str) else _num(x)
                    if p.get('precio') is not None:
                        p['precio'] = round(p['precio'] / 1000.0, 3)     # miles -> mdp per unit
                    protos.append(p)
    ini, fin = fechas.get('fecha de inicio') or [], fechas.get('fecha de termino') or []
    out['fechas'] = {'proforma': [ini[0] if ini else None, fin[0] if fin else None],
                     'flujos': [ini[1] if len(ini) > 1 else None, fin[1] if len(fin) > 1 else None]}
    out['u'] = unidades[0] if unidades else None
    out['u_tipo'] = unidades[1] if unidades else None
    out['prototipos'] = [p for p in protos if p.get('n')]

    # funding: column B labels after the P&L, amount in column C; Ranman's
    # own contributions and returns in column H/I
    fondeo = {}
    for i in range(fin_pyl + 1, b):
        r = filas[i]
        k = _FONDEO.get(_norm(r[1]) if len(r) > 1 else '')
        if k and k not in fondeo:
            fondeo[k] = _mdp(r[2]) if len(r) > 2 else None
        for j in range(4, min(len(r) - 1, 12)):
            t = _norm(r[j]) if isinstance(r[j], str) else ''
            if t in ('aportaciones', 'recuperaciones', 'tir', 'moic') and t not in fondeo:
                v = _num(r[j + 1])
                if t in ('aportaciones', 'recuperaciones'):
                    fondeo[t] = round(v / 1000.0, 3) if v is not None else None
                elif t == 'tir':
                    fondeo[t] = round(v, 4) if v is not None and -1 < v < 5 else None
                else:
                    fondeo[t] = round(v, 3) if v is not None and 0 < v < 50 else None
    out['fondeo'] = fondeo
    return out


def _vacio(e):
    return all(not v for s in e['esc'].values() for v in s.values())


def _cuadra(s, tol):
    """The P&L identities of one column -> [(line, difference)] that miss."""
    if s.get('ventas') is None:
        return []
    v = {k: (s.get(k) or 0.0) for k, _ in CONCEPTOS}
    checks = [('margen_bruto', v['ventas'] - sum(v[k] for k in COSTOS)),
              ('ebit', v['margen_bruto'] - sum(v[k] for k in GASTOS)),
              ('uair', v['ebit'] - sum(v[k] for k in FINANCIEROS))]
    return [(k, round(v[k] - calc, 3)) for k, calc in checks if abs(v[k] - calc) > tol]


# ── «Saldos Finales»: month-end cash by etapa ───────────────────────────────
_PRECIO = re.compile(r'\(\s*([\d.]+)(?:\s*a\s*[\d.]+)?\s*mdp\s*\)\s*', re.I)


def _saldos(ws):
    """{meses, filas:[{n, k, ticket, v}]} in mdp, from the first block of etapa
    rows (the subtotals and the plaza total below it are left out)."""
    filas = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 80), values_only=True))
    hm = next((i for i, r in enumerate(filas[:8])
               if sum(1 for v in r if isinstance(v, (dt.datetime, dt.date))) >= 12), None)
    if hm is None:
        return None
    cm = [(j, _mes(v)) for j, v in enumerate(filas[hm]) if isinstance(v, (dt.datetime, dt.date))]
    j0 = cm[0][0]
    hp = next((i for i in range(hm, min(hm + 6, len(filas)))
               if any(_norm(v).startswith('proyecto') for v in filas[i][:j0])), None)
    if hp is None:
        return None
    jn = next(j for j, v in enumerate(filas[hp][:j0]) if _norm(v).startswith('proyecto'))
    out, previa, acum = [], None, []
    for r in filas[hp + 1:]:
        nom = r[jn] if jn < len(r) else None
        if not isinstance(nom, str) or not nom.strip():
            if out:
                break
            continue
        t = _norm(nom)
        vals = [_mdp(r[j]) if j < len(r) else None for j, _ in cm]
        vals = [v if v is not None and abs(v) > 0.0005 else 0.0 for v in vals]
        if t.startswith(('total', 'subtotal', 'saldo')):
            acum = []
            continue
        # a row that is the sum of the rows above it (Celaya's «Miranda») is a subtotal
        if len(acum) >= 2 and any(vals) and all(abs(v - sum(x[m] for x in acum)) < 0.01 for m, v in enumerate(vals)):
            acum = []
            continue
        m = _PRECIO.search(nom)
        limpio = _PRECIO.sub('', nom).strip()
        k = clave_desarrollo(limpio) or previa
        previa = k
        out.append({'n': limpio, 'k': k, 'ticket': float(m.group(1)) if m else None, 'v': vals})
        acum.append(vals)
    if not out:
        return None
    usados = [m for m in range(len(cm)) if any(f['v'][m] for f in out)]
    if not usados:
        return {'meses': [], 'filas': [dict(f, v=[]) for f in out]}
    a, b = usados[0], usados[-1] + 1
    return {'meses': [ym for _, ym in cm[a:b]], 'filas': [dict(f, v=f['v'][a:b]) for f in out]}


def parse_bp(file_bytes: bytes, filename: str, as_of: str = None):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = _hoja(wb, 'supuestos')
        if ws is None:
            raise SinSupuestos("%s has no 'Supuestos' sheet (one block per etapa), so there is nothing to read" % filename)
        filas = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=34, values_only=True))
        sf = _hoja(wb, 'saldos finales')
        saldos = _saldos(sf) if sf is not None else None
    finally:
        wb.close()

    corte = None
    for r in filas[:4]:
        for j, v in enumerate(r):
            if _norm(v).startswith('informacion historica'):
                corte = _mes(next((x for x in r[j + 1:j + 6] if isinstance(x, (dt.datetime, dt.date))), None))
    if not corte:
        raise ValueError("%s: no «INFORMACIÓN HISTÓRICA» month on the Supuestos sheet — can't tell the cut" % filename)
    bloques = _bloques(filas)
    if not bloques:
        raise ValueError("%s: no etapa blocks on the Supuestos sheet" % filename)
    etapas = [e for e in (_lee_bloque(filas, a, b) for a, b in bloques) if not _vacio(e)]

    ok, lines = True, []
    for e in etapas:
        for esc, et in ESCENARIOS:
            s = e['esc'][esc]
            tol = max(0.01, 0.002 * abs(s.get('ventas') or 0))
            fallas = _cuadra(s, tol)
            if fallas and esc != 'proforma':          # the signed proforma is typed in by hand
                ok = False
            for k, d in fallas:
                lines.append('%-34s %-16s %-14s no cuadra: dif %.3f mdp%s'
                             % (e['n'][:34], et, k, d, ' (aviso)' if esc == 'proforma' else ''))
    if not etapas:
        raise ValueError("%s: every etapa block on the Supuestos sheet is empty" % filename)
    lines.append('%d etapas: P&L cuadra en cierre previsto y flujos' % len(etapas) if ok else
                 'no se guardó: hay etapas cuyo P&L no cuadra')

    data = {'archivo': filename, 'modelo': nombre_de(filename), 'slug': slug_de(filename),
            'corte': corte, 'asOf': _fin_de_mes(corte), 'conceptos': [[k, et] for k, et in CONCEPTOS],
            'escenarios': [[k, et] for k, et in ESCENARIOS], 'etapas': etapas, 'saldos': saldos}
    return data, bool(ok), lines
