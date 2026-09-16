"""RANMAN «Flujo y PLP DRA - MAQUINA - dd-mm-aa.xlsx» -> the Flujo y PLP tablero's
flujo_plp.json shape.

Faithful port of Valoran's actualizar_flujo_plp.py (Sep-2026), working from
bytes so the workbook can be uploaded through the app. Also validates a
flujo_plp.json produced by the original script.

Where each thing comes from (sheet «Resumen» is in thousands of pesos; the
tablero shows mdp):
    Resumen  rows  7-13    aplicación · operativos
                   17-24   aplicación · deuda y proyectos no operativos
                   29-41   aplicación · inversión
                   47-57   origen
                   44,58,60 control totals (aplicación, origen, Flujo Ranman)
                   79-86   corporate debt balances
             col A        1/0 flag that switches each row on or off
    PLP      rows  4-82    income statement and cash flow
             cols  C:EP    monthly block (144 months)
                   KG:KR   yearly block (12 years)

Every Resumen line is re-computed as if its flag were 1 ("base" amount) so the
tablero can switch on lines the file has off. The validation then checks that,
with the flags as they come, the recomputed totals match the file's own.
"""
import datetime as dt
import io
import json
import re

import openpyxl

GRUPOS = [('op', 'Operativos',                      range(7, 14)),
          ('de', 'Deuda y proyectos no operativos', range(17, 25)),
          ('in', 'Inversión',                       range(29, 42)),
          ('or', 'Origen',                          range(47, 58))]
FILA_APLICACION, FILA_ORIGEN, FILA_FLUJO = 44, 58, 60
FILAS_DEUDA, FILA_SALDO = range(79, 86), 86
PARAM = {89, 90}          # $A$89 / $A$90 are sale percentages, not flags
MILES = 1000.0
RENOMBRES = {'Crédito Valoran': 'Crédito VRM'}   # how each line is called in the tablero

PLP_R0, PLP_R1 = 4, 82
COLS_MES = range(3, 147)      # C .. EP
COLS_ANIO = range(293, 305)   # KG .. KR
FILAS_CUENTA = {4, 5}         # apartados y firmas: units, not thousands


def fecha_de_nombre(nombre: str):
    """'Flujo y PLP DRA - MAQUINA - 31-08-26.xlsx' -> '2026-08-31' (None if absent)."""
    m = re.search(r'(\d{2})-(\d{2})-(\d{2})\.xlsx$', nombre or '')
    if not m:
        return None
    d, mes, a = (int(x) for x in m.groups())
    try:
        return dt.date(2000 + a, mes, d).isoformat()
    except ValueError:
        return None


# ── sheet «Resumen»: origen, aplicación y deuda ─────────────────────────────
def lee_resumen(wf, wv, as_of: str) -> dict:
    sf, sv, flv = wf['Resumen'], wv['Resumen'], wv['Flujo']

    # row 1 carries the monthly block and then yearly columns: cut where dates go back
    hdr = list(sv.iter_rows(min_row=1, max_row=1, max_col=133, values_only=True))[0]
    fe = [(j, c) for j, c in enumerate(hdr) if isinstance(c, dt.datetime)]
    for k in range(1, len(fe)):
        if fe[k][1] < fe[k - 1][1]:
            fe = fe[:k]
            break
    COLS = [j + 1 for j, _ in fe]
    MESES = [c.strftime('%Y-%m') for _, c in fe]

    # the series starts at the corte's month: earlier columns come as zeros
    corte = as_of[:7]
    if MESES and MESES[0] < corte and corte in MESES:
        k = MESES.index(corte)
        COLS, MESES = COLS[k:], MESES[k:]

    LINEAS = [r for _, _, rs in GRUPOS for r in rs]
    memo = {}
    num = lambda v: v if isinstance(v, (int, float)) else 0.0

    def base(r, c, prof=0):
        """The row's value as if its flag were 1."""
        if prof > 12:
            return 0.0
        k = (r, c)
        if k in memo:
            return memo[k]
        memo[k] = 0.0
        f = sf.cell(r, c).value
        if f is None:
            v = num(sv.cell(r, c).value)
        elif isinstance(f, (int, float)):
            v = float(f)
        else:
            s = str(f).strip().lstrip('=').lstrip('+')
            m = re.match(r'^IF\(\$A\$\d+=1,([^,]+),', s)      # the flag as a mode selector
            if m:
                s = m.group(1)
            s = re.sub(r'Flujo!\$?([A-Z]{1,3})\$?(\d+)',
                       lambda m: repr(num(flv['%s%s' % (m.group(1), m.group(2))].value)), s)
            s = re.sub(r'\$A\$(\d+)',
                       lambda m: repr(num(sv.cell(int(m.group(1)), 1).value))
                       if int(m.group(1)) in PARAM else '1', s)

            def ref(m):
                col = openpyxl.utils.column_index_from_string(m.group(1))
                row = int(m.group(2))
                return repr(base(row, col, prof + 1) if row in LINEAS
                            else num(sv.cell(row, col).value))
            s = re.sub(r'\$?([A-Z]{1,3})\$?(\d+)', ref, s).replace('%', '/100')
            try:
                v = float(eval(s, {'__builtins__': {}}, {}))
            except Exception:
                v = num(sv.cell(r, c).value)
        memo[k] = v
        return v

    fila = lambda r: [round(num(sv.cell(r, c).value) / MILES, 2) for c in COLS]

    lineas = []
    for cod, _, rs in GRUPOS:
        for r in rs:
            lab = ' '.join(str(sv.cell(r, c).value) for c in range(2, 7)
                           if sv.cell(r, c).value is not None).strip()
            if not lab:
                continue
            serie = [round(base(r, c) / MILES, 2) for c in COLS]
            lineas.append({'r': r, 'g': cod, 'n': lab,
                           'on': bool(sv.cell(r, 1).value not in (0, None)),
                           'v': serie,
                           'hay': any(abs(x) > 0.005 for x in serie)})

    deuda = {}
    for r in FILAS_DEUDA:
        nom = sv.cell(r, 2).value or sv.cell(r, 3).value or sv.cell(r, 4).value
        if nom:
            nom = str(nom).strip()
            deuda[RENOMBRES.get(nom, nom)] = fila(r)

    return {'meses': MESES, 'lineas': lineas, 'deuda': deuda,
            'saldoNec': fila(FILA_SALDO),
            'ref': {'aplic': fila(FILA_APLICACION), 'origen': fila(FILA_ORIGEN),
                    'flujo': fila(FILA_FLUJO)}}


# ── sheet «PLP»: income statement and cash flow, rows 4 to 82 ───────────────
def lee_plp(wv, as_of: str) -> dict:
    ws = wv['PLP']
    meses = [ws.cell(3, c).value for c in COLS_MES]
    anios = [ws.cell(3, c).value for c in COLS_ANIO]
    if not all(isinstance(m, dt.datetime) for m in meses):
        raise ValueError('La hoja PLP no trae fechas en C3:EP3. ¿Se movieron las columnas?')

    # the tablero only shows the current year and the next
    a0 = int(as_of[:4])
    idx = [i for i, m in enumerate(meses) if m.year in (a0, a0 + 1)]
    num = lambda v: v if isinstance(v, (int, float)) else None

    filas = []
    for r in range(PLP_R0, PLP_R1 + 1):
        et = ws.cell(r, 2).value
        if et is None:
            filas.append({'r': r, 'sep': 1})
            continue
        fmt = ws.cell(r, 3).number_format or ''
        t = 'pct' if '%' in fmt else ('cuenta' if r in FILAS_CUENTA else 'mdp')

        def val(c, t=t, r=r):
            v = num(ws.cell(r, c).value)
            if v is None:
                return None
            return round(v / 1000, 2) if t == 'mdp' else (round(v, 4) if t == 'pct' else round(v))
        f = {'r': r, 'n': str(et).strip(), 'niv': ws.row_dimensions[r].outline_level or 0, 't': t,
             'a': [val(c) for c in COLS_ANIO],
             'm': [val(list(COLS_MES)[i]) for i in idx]}
        if ws.row_dimensions[r].hidden:
            f['plg'] = 1
        filas.append(f)

    return {'anios': [int(a) for a in anios],
            'meses': [meses[i].strftime('%Y-%m') for i in idx],
            'filas': filas}


# ── validation: recomputed totals vs the file's own ─────────────────────────
def valida(fl: dict, plp: dict):
    """(ok, lines) — every check as a human-readable line, like the script prints."""
    ok, out = True, []
    n = len(fl['meses'])
    grupos = {'aplic': {'op', 'de', 'in'}, 'origen': {'or'}}
    for clave, gs in grupos.items():
        calc = [sum(l['v'][i] for l in fl['lineas'] if l['g'] in gs and l['on']) for i in range(n)]
        dif = max(abs(a - b) for a, b in zip(calc, fl['ref'][clave])) if n else 0.0
        ok &= dif < 0.1
        out.append('%-14s dif máx contra el archivo: %6.2f mdp' % (clave, dif))
    calc = [sum(l['v'][i] for l in fl['lineas'] if l['on']) for i in range(n)]
    dif = max(abs(a - b) for a, b in zip(calc, fl['ref']['flujo'])) if n else 0.0
    ok &= dif < 0.1
    out.append('%-14s dif máx contra el archivo: %6.2f mdp' % ('flujo Ranman', dif))

    fs = plp['filas']
    for i, f in enumerate(fs):
        if f.get('sep') or f.get('niv', 1) != 0:
            continue
        h = []
        for j in range(i + 1, len(fs)):
            g = fs[j]
            if g.get('sep') or g.get('niv', 0) == 0:
                break
            h.append(g)
        if not h:
            continue
        sub = [x for x in h if x['n'] in ('Administración', 'Comercialización')] or h
        dif = max(abs((f['a'][k] or 0) - sum(x['a'][k] or 0 for x in sub))
                  for k in range(len(plp['anios'])))
        ok &= dif < 0.1
        out.append('PLP renglón %-3d %-42s dif máx %6.2f mdp' % (f['r'], f['n'][:42], dif))
    return bool(ok), out


# ── entry points ────────────────────────────────────────────────────────────
def parse_workbook(file_bytes: bytes, filename: str):
    """One «Flujo y PLP DRA - MAQUINA - dd-mm-aa.xlsx» -> (datos, ok, lines).

    datos has the tablero's flujo_plp.json shape. Raises ValueError when the
    date can't be read from the name or the sheets aren't there."""
    as_of = fecha_de_nombre(filename)
    if not as_of:
        raise ValueError("%s: could not read the date from the file name (expected '... dd-mm-aa.xlsx')" % filename)
    wf = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)   # formulas
    wv = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)    # cached values
    try:
        for hoja in ('Resumen', 'Flujo', 'PLP'):
            if hoja not in wv.sheetnames:
                raise ValueError("%s has no '%s' sheet — is it a Flujo y PLP DRA workbook?" % (filename, hoja))
        fl = lee_resumen(wf, wv, as_of)
        plp = lee_plp(wv, as_of)
    finally:
        wf.close()
        wv.close()
    ok, lines = valida(fl, plp)
    datos = {'archivo': filename, 'asOf': as_of,
             'generado': dt.datetime.now().isoformat(timespec='seconds'),
             'flujo': fl, 'plp': plp}
    return datos, ok, lines


def parse_json(file_bytes: bytes) -> dict:
    """A flujo_plp.json written by the original script; validated for shape."""
    data = json.loads(file_bytes.decode('utf-8'))
    if not isinstance(data, dict) or not data.get('asOf') or not isinstance(data.get('flujo'), dict) \
            or not isinstance(data.get('plp'), dict):
        raise ValueError('flujo_plp.json should carry asOf, flujo and plp.')
    dt.date.fromisoformat(data['asOf'])
    fl, plp = data['flujo'], data['plp']
    for k in ('meses', 'lineas', 'deuda', 'saldoNec', 'ref'):
        if k not in fl:
            raise ValueError("flujo_plp.json: flujo is missing '%s'." % k)
    for k in ('anios', 'meses', 'filas'):
        if k not in plp:
            raise ValueError("flujo_plp.json: plp is missing '%s'." % k)
    if not fl['meses'] or not fl['lineas'] or not plp['filas']:
        raise ValueError('flujo_plp.json has no data in it.')
    return data
