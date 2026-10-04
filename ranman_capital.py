"""RANMAN's cash, capital and growth files -> data for the Finance, Cashflow &
PLP and Strategy tabs.

    parse_ccc           «Control de CCC al dd-mm-aa.xlsx»            the Bancrea 48 and BBVA corporate lines
    parse_aportaciones  «Aportaciones Maquina al dd-mm-aa.xlsx»      Maquina's contributions, 2023 on
    parse_prestamos     «Control de préstamos entre proyectos …»     intercompany loans, lender by borrower
    parse_necesidad     «Necesidad de Capital al aa-mm-dd.xlsx»      this week's payments against cash, by project
    parse_reinversion   «Reinversión de utilidades.xlsx»             when each coto frees equity, three versions
    parse_reservas      «TABLA RESERVAS TERRITORIALES AT.xlsx»       the land pipeline

Every parser returns (data, ok, lines) like the others in ranman_package: ok
False means the file's own totals don't add up and nothing should be stored.
Nothing is read by fixed rows; tables are found by their headers. Amounts come
in pesos (CCC, Aportaciones, Préstamos, Necesidad) or thousands (Reinversión)
and are stored in mdp.
"""
import datetime as dt
import io
import re

import openpyxl

from ranman_package import _norm, _num, _hoja, clave_desarrollo


def _mdp(v, escala=1e6):
    """Pesos -> mdp, kept to the peso (6 decimals) so the files' own sums still close."""
    n = _num(v)
    return round(n / escala, 6) if n is not None else None


def _fecha(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else (v.isoformat() if isinstance(v, dt.date) else None)


def _filas(ws, max_col=40, max_row=None):
    return list(ws.iter_rows(min_row=1, max_row=max_row or ws.max_row, max_col=max_col, values_only=True))


def _celda(filas, i, j):
    return filas[i][j] if 0 <= i < len(filas) and j < len(filas[i]) else None


def _busca(filas, pred, filas_max=None, desde=0):
    """(row, col) of the first cell whose normalised text satisfies pred."""
    for i in range(desde, min(len(filas), filas_max or len(filas))):
        for j, v in enumerate(filas[i]):
            if isinstance(v, str) and pred(_norm(v)):
                return i, j
    return None


def _cuadra(a, b, tol=1.0):
    return a is not None and b is not None and abs(a - b) <= tol


# ── «Control de CCC al dd-mm-aa.xlsx» ───────────────────────────────────────
# One sheet, «Control», six tables side by side under titles in row 1: the two
# ledgers of what each draw paid for (A, D), the dated draws and repayments of
# each line (G: «Disposiciones Bancrea CCC 48 MDP», then «Disposiciones BBVA
# CCC 58 MDP» further down, sometimes repeated month by month), Plan Houston
# (plan vs drawn by project, K), the balance by project (O) and by year (T).
# A line's movements add up to the «Dispuesto» or the «Por disponer» it
# declares — Bancrea counts draws as positive, the BBVA block counts the
# repayments («Pago al CCC (Ruba)») as positive — and that is the check.
_PAGO = re.compile(r'pago|devoluci|amortiz')


def _bloque_linea(filas, i0, j):
    """The dated movements and summary rows of one line's block, from its title row."""
    titulo = str(filas[i0][j]).strip()
    m = re.search(r'(\d+(?:\.\d+)?)\s*md', _norm(titulo))
    limite = float(m.group(1)) if m else None
    movs, resumen, fin = [], [], len(filas)
    for i in range(i0 + 1, len(filas)):
        f, c, v = _celda(filas, i, j), _celda(filas, i, j + 1), _celda(filas, i, j + 2)
        t = _norm(f) if isinstance(f, str) else ''
        if t.startswith('disposiciones '):
            fin = i
            break
        if t in ('fecha',):
            continue
        if isinstance(f, (dt.datetime, dt.date)) and _num(v) is not None:
            et = str(c or '').strip()
            movs.append({'f': _fecha(f), 'n': et, 'v': round(_num(v), 2), 'pago': bool(_PAGO.search(_norm(et))),
                         'devolucion': 'devoluci' in _norm(et)})
        elif t and _num(v) is not None:
            resumen.append((str(f).strip(), round(_num(v), 2)))
        elif not t and f is None and c is None and v is None and movs and resumen:
            # two blank rows after the summary close the block
            if all(_celda(filas, k, j) is None for k in range(i, min(i + 2, len(filas)))):
                fin = i
                break
    return {'titulo': titulo, 'limite': limite, 'movs': movs, 'resumen': resumen}, fin


def _tabla_cols(filas, i0, j0, ancho):
    """Rows under a block title: [name, values...] until a «Total» row (included)."""
    out = []
    for i in range(i0 + 2, len(filas)):
        n = _celda(filas, i, j0)
        if n is None or (isinstance(n, str) and not n.strip()):
            if out and any(_norm(r[0]).startswith('total') for r in out):
                break
            continue
        vals = [_num(_celda(filas, i, j0 + k)) for k in range(1, ancho + 1)]
        out.append([str(n).strip()] + vals)
        if _norm(n).startswith('total') and not _norm(n).startswith('total ccc'):
            if ancho and all(v is None for v in vals):
                continue
            break
        if _norm(n).startswith('total ccc'):
            break
    return out


def parse_ccc(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = _hoja(wb, 'control') or wb.worksheets[0]
        filas = _filas(ws, max_col=30)
    finally:
        wb.close()
    ok, lines = True, []
    # the lines: every «Disposiciones …» block in its column, first occurrence of each title
    pos = _busca(filas, lambda t: t.startswith('disposiciones '), 4)
    if pos is None:
        raise ValueError("%s: no «Disposiciones …» table on the Control sheet" % filename)
    jd = pos[1]
    lineas, vistos, i = [], set(), pos[0]
    while i < len(filas):
        v = _celda(filas, i, jd)
        if isinstance(v, str) and _norm(v).startswith('disposiciones '):
            b, fin = _bloque_linea(filas, i, jd)
            clave = re.sub(r'\d+(\.\d+)?\s*md.*$', '', _norm(b['titulo'])).strip()
            if clave not in vistos and b['movs']:
                vistos.add(clave)
                lineas.append(b)
            i = max(fin, i + 1)
        else:
            i += 1
    for b in lineas:
        s = round(sum(m['v'] for m in b['movs']), 2)
        decl = {_norm(k): v for k, v in b['resumen']}
        dispuesto, por_disp = decl.get('dispuesto'), decl.get('por disponer')
        cuadra = _cuadra(s, dispuesto) or _cuadra(s, por_disp) or (
            b['limite'] and por_disp is not None and _cuadra(b['limite'] * 1e6 - s, por_disp))
        ok &= bool(cuadra)
        lines.append('%-36s movimientos %.2f vs dispuesto %s / por disponer %s: dif %.2f'
                     % (b['titulo'][:36], s, dispuesto, por_disp,
                        0.0 if cuadra else min(abs(s - (dispuesto or 0)), abs(s - (por_disp or 0)))))
        draws = sum(abs(m['v']) for m in b['movs'] if not m['pago'])
        # money handed back to the line («Devolución») nets the draws; the rest
        # of the payments («Pago al CCC (Ruba)») come from asset sales
        devuelto = sum(abs(m['v']) for m in b['movs'] if m['devolucion'])
        pagos = sum(abs(m['v']) for m in b['movs'] if m['pago'] and not m['devolucion'])
        # the payments that name the sale behind them («Pago al CCC (Kingfa)»)
        ventas_act = sum(abs(m['v']) for m in b['movs'] if m['pago'] and not m['devolucion'] and '(' in m['n'])
        reservas = [(k, v) for k, v in b['resumen'] if _norm(k) not in ('dispuesto', 'por disponer', 'disponible')]
        disponible = decl.get('disponible')
        if disponible is None and por_disp is not None:
            disponible = por_disp - sum(v for _, v in reservas)
        b.update(dispuesto=round((draws - devuelto) / 1e6, 6), devuelto=round(devuelto / 1e6, 6), pagado=round(pagos / 1e6, 6),
                 pagado_ventas=round(ventas_act / 1e6, 6),
                 por_disponer=_mdp(por_disp) if por_disp is not None else (
                     round(b['limite'] - s / 1e6, 6) if b['limite'] else None),
                 reservado=[{'n': k, 'v': round(v / 1e6, 6)} for k, v in reservas],
                 disponible=_mdp(disponible),
                 movs=[dict(m, v=round(m['v'] / 1e6, 6)) for m in b['movs']])
        b['nombre'] = 'Bancrea' if 'bancrea' in _norm(b['titulo']) else ('BBVA' if 'bbva' in _norm(b['titulo']) else b['titulo'])
        del b['resumen']

    # the limit: the ledgers' titles («Ministraciones BBVA CCC 55 MDP») agree with
    # the Cuadro de Riesgos where the draws table's title does not («… 58 MDP»)
    for v in (filas[0] if filas else []):
        t = _norm(v) if isinstance(v, str) else ''
        m = re.search(r'(\d+(?:\.\d+)?)\s*m', t) if t.startswith('ministraciones') else None
        for b in lineas:
            if m and b['nombre'].lower() in t:
                lim = float(m.group(1))
                if b['limite'] and abs(b['limite'] - lim) > 0.01:
                    b['limite_otro'] = b['limite']
                b['limite'] = lim

    tablas = {}
    for clave, pred, ancho in (('plan', lambda t: t.startswith('plan houston'), 2),
                               ('proyectos', lambda t: t.startswith('saldos por proyecto'), 3),
                               ('anios', lambda t: t.startswith('saldos por a'), 3)):
        p = _busca(filas, pred, 3)
        if p is None:
            continue
        cab = [str(_celda(filas, p[0] + 1, p[1] + k) or '').strip() for k in range(0, ancho + 1)]
        tablas[clave] = {'titulo': str(filas[p[0]][p[1]]).strip(), 'cols': cab[1:],
                         'filas': _tabla_cols(filas, p[0], p[1], ancho)}
    # balances by project and by year add up to their Total row
    for clave in ('proyectos', 'anios'):
        t = tablas.get(clave)
        if not t or not t['filas']:
            continue
        tot = next((r for r in t['filas'] if _norm(r[0]) == 'total'), None)
        cuerpo = [r for r in t['filas'] if r is not tot]
        if tot:
            for k in range(1, len(tot)):
                s = sum(r[k] or 0 for r in cuerpo)
                c = _cuadra(s, tot[k], 2.0)
                ok &= c or tot[k] is None
                lines.append('Saldos por %-9s %-16s suma %.2f vs Total %s: dif %.2f'
                             % ('proyecto' if clave == 'proyectos' else 'año', t['cols'][k - 1][:16], s, tot[k], abs(s - (tot[k] or 0))))
    # Plan Houston: each «Subtotal» closes the projects above it, «Total CCC» the subtotals
    plan = tablas.get('plan')
    if plan:
        grupo, grupos = [], []
        for r in plan['filas']:
            t = _norm(r[0])
            if t.startswith('subtotal'):
                s = sum(x[2] or 0 for x in grupo)
                ok &= _cuadra(s, r[2], 2.0)
                lines.append('Plan Houston %-22s dispuesto %.2f vs subtotal %s: dif %.2f' % (r[0][:22], s, r[2], abs(s - (r[2] or 0))))
                grupos.append({'n': r[0], 'plan': _mdp(r[1]), 'dispuesto': _mdp(r[2]),
                               'filas': [{'n': x[0], 'k': clave_desarrollo(x[0]), 'plan': _mdp(x[1]), 'dispuesto': _mdp(x[2])} for x in grupo]})
                grupo = []
            elif t.startswith('total'):
                plan['total'] = {'plan': _mdp(r[1]), 'dispuesto': _mdp(r[2])}
            else:
                grupo.append(r)
        plan['grupos'] = grupos
        del plan['filas']
    for clave in ('proyectos', 'anios'):
        t = tablas.get(clave)
        if t:
            t['filas'] = [{'n': r[0], 'k': clave_desarrollo(r[0]), 'v': [_mdp(x) for x in r[1:]], 'total': _norm(r[0]) == 'total'}
                          for r in t['filas']]
    if not lineas:
        raise ValueError("%s: no draws found under the «Disposiciones …» titles" % filename)
    data = {'archivo': filename, 'asOf': as_of, 'lineas': lineas, **tablas}
    return data, bool(ok), lines


# ── «Aportaciones Maquina al dd-mm-aa.xlsx» ─────────────────────────────────
# Sheet «Total»: one row per contribution (number, month) with its origin
# (Maquina MX / USA / Casas Muestra), what it paid for (land, operations, bridge
# and long-term interest, corporate interest) and the project, each group closed
# by its own «Total» column; then a Total row and Ranman's «Validación» row.
# The hidden «Préstamos» sheet lists Maquina's loans to Ranman, repayable on
# liquidity events; the hidden «Resumen» groups the uses (operativos, deuda y
# proyectos no operativos, inversión).
def _grupos_cab(filas, i2, i3):
    """{grupo: [(col, label)]} from the two header rows; each group runs up to its «Total» column."""
    grupos, actual = {}, None
    for j in range(len(filas[i3])):
        g = _celda(filas, i2, j)
        if isinstance(g, str) and g.strip():
            actual = _norm(g)
            grupos[actual] = []
        lab = _celda(filas, i3, j)
        if actual and isinstance(lab, str) and lab.strip():
            grupos[actual].append((j, lab.strip()))
            if _norm(lab) == 'total':
                actual = None
    return grupos


def parse_aportaciones(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = _hoja(wb, 'total')
        if ws is None:
            raise ValueError("%s has no 'Total' sheet" % filename)
        filas = _filas(ws, max_col=40)
        pr = _hoja(wb, 'prestamos')
        prest = _filas(pr, max_col=4) if pr is not None else []
        rs = _hoja(wb, 'resumen')
        resumen = _filas(rs, max_col=6) if rs is not None else []
    finally:
        wb.close()
    p = _busca(filas, lambda t: t == 'origen', 6)
    if p is None:
        raise ValueError("%s: no «Origen» header on the Total sheet" % filename)
    i2, i3 = p[0], p[0] + 1
    grupos = _grupos_cab(filas, i2, i3)
    g_or = next((v for k, v in grupos.items() if k.startswith('origen')), [])
    g_ti = next((v for k, v in grupos.items() if 'tipo' in k), [])
    g_pr = next((v for k, v in grupos.items() if 'proyecto' in k), [])
    if not (g_or and g_ti and g_pr):
        raise ValueError("%s: the Origen / Destino (Tipo) / Destino (Proyecto) headers are not all there" % filename)
    jm = next((j for j, v in enumerate(filas[i2]) if _norm(v) == 'mes'), 1)
    sin_total = lambda g: [(j, l) for j, l in g if _norm(l) != 'total']
    col_total = lambda g: next((j for j, l in g if _norm(l) == 'total'), None)
    filas_ap, tot_row, val_row = [], None, None
    for r in filas[i3 + 1:]:
        etq = _norm(r[jm]) if jm < len(r) and isinstance(r[jm], str) else ''
        if etq == 'total':
            tot_row = r
            continue
        if etq.startswith('validaci'):
            val_row = r
            break
        if tot_row is not None:
            continue
        f = r[jm] if jm < len(r) else None
        if not isinstance(f, (dt.datetime, dt.date)):
            continue
        filas_ap.append({'n': _num(r[0]), 'f': _fecha(f),
                         'origen': {l: _mdp(r[j]) or 0.0 for j, l in sin_total(g_or)},
                         'tipo': {l: _mdp(r[j]) or 0.0 for j, l in sin_total(g_ti)},
                         'proyecto': {l: _mdp(r[j]) or 0.0 for j, l in sin_total(g_pr) if _mdp(r[j])},
                         't': _mdp(r[col_total(g_or)]) if col_total(g_or) is not None else None})
    if not filas_ap or tot_row is None:
        raise ValueError("%s: no contributions or no Total row on the Total sheet" % filename)
    ok, lines = True, []
    malas = 0
    for a in filas_ap:
        so, st, sp = sum(a['origen'].values()), sum(a['tipo'].values()), sum(a['proyecto'].values())
        if not (abs(so - st) < 2e-5 and abs(so - sp) < 2e-5):
            malas += 1
            lines.append('Aportación %s (%s): origen %.4f, tipo %.4f, proyecto %.4f — dif %.4f mdp'
                         % (a['n'], a['f'], so, st, sp, max(abs(so - st), abs(so - sp))))
    ok &= malas == 0
    for g, nom in ((g_or, 'origen'), (g_ti, 'tipo'), (g_pr, 'proyecto')):
        jt = col_total(g)
        decl = _mdp(tot_row[jt]) if jt is not None else None
        s = sum(sum(a[nom].values()) for a in filas_ap)
        c = _cuadra(s, decl, 2e-4)
        ok &= c
        lines.append('Total %-9s suma %.4f vs fila Total %s: dif %.6f mdp' % (nom, s, decl, abs(s - (decl or 0))))
    if val_row is not None:
        peor = max((abs(_num(v) or 0) for v in val_row if _num(v) is not None), default=0)
        ok &= peor < 1.0
        lines.append('Fila Validación del archivo: mayor dif %.2f pesos' % peor)
    # Maquina's loans to Ranman (hidden sheet)
    prestamos, prest_total = [], None
    for r in prest:
        if isinstance(r[0], (dt.datetime, dt.date)) and _num(r[1]) is not None:
            prestamos.append({'f': _fecha(r[0]), 'v': _mdp(r[1]), 'uso': str(r[2] or '').strip()})
        elif isinstance(r[0], str) and _norm(r[0]) == 'total':
            prest_total = _mdp(r[1])
    if prest_total is not None:
        s = sum(x['v'] for x in prestamos)
        ok &= _cuadra(s, prest_total, 1e-5)
        lines.append('Préstamos MAQUINA: suma %.4f vs Total %.4f: dif %.6f mdp' % (s, prest_total, abs(s - prest_total)))
    # what it went to, as Ranman groups it (hidden «Resumen»)
    aplicacion, grupo = [], None
    for r in resumen:
        g, n, v = (r[1] if len(r) > 1 else None), (r[2] if len(r) > 2 else None), (_num(r[3]) if len(r) > 3 else None)
        if isinstance(g, str) and g.strip() and not (isinstance(n, str) and n.strip()):
            grupo = {'n': g.strip(), 'filas': [], 'total': None}
            aplicacion.append(grupo)
        if grupo and isinstance(n, str) and n.strip() and v is not None:
            if _norm(n).startswith('total aplic'):
                break
            if _norm(n).startswith('total'):
                grupo['total'] = round(v / 1e6, 4)
            else:
                grupo['filas'].append({'n': n.strip(), 'k': clave_desarrollo(n), 'v': round(v / 1e6, 4)})
    data = {'archivo': filename, 'asOf': as_of, 'aportaciones': filas_ap,
            'origenes': [l for _, l in sin_total(g_or)], 'tipos': [l for _, l in sin_total(g_ti)],
            'proyectos': [{'n': l, 'k': clave_desarrollo(l)} for _, l in sin_total(g_pr)],
            'total': _mdp(tot_row[col_total(g_or)]), 'prestamos': prestamos, 'prestamos_total': prest_total,
            'aplicacion': [g for g in aplicacion if g['filas']]}
    return data, bool(ok), lines


# ── «Control de préstamos entre proyectos … (dd-mm-aa)).xlsx» ───────────────
# «Resumen»: who lent (columns, «Cuentas por Cobrar») to whom (rows, «Cuentas
# por Pagar»), with a Total column and row. A row followed straight away by
# rows that add up to it («Corporativo (Proyectos no operativos)» over PP GDL,
# Las Veredas, PP QRO) is a group. «Integración» is the dated ledger behind it;
# the hidden «Otros» holds loans already settled (Miranda's 4.5 mdp).
def _ledger(filas):
    movs, total = [], None
    for r in filas:
        if isinstance(r[0], (dt.datetime, dt.date)) and len(r) > 3 and _num(r[3]) is not None:
            movs.append({'f': _fecha(r[0]), 'presta': str(r[1] or '').strip(), 'recibe': str(r[2] or '').strip(),
                         'v': _mdp(r[3])})
        elif any(isinstance(x, str) and _norm(x) == 'total' for x in r[:3]) and len(r) > 3 and _num(r[3]) is not None:
            total = _mdp(r[3])
    return movs, total


def parse_prestamos(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        rs = _hoja(wb, 'resumen')
        if rs is None:
            raise ValueError("%s has no 'Resumen' sheet" % filename)
        filas = _filas(rs, max_col=30)
        it = _hoja(wb, 'integracion')
        integ = _filas(it, max_col=5) if it is not None else []
        ot = _hoja(wb, 'otros')
        otros = _filas(ot, max_col=5) if ot is not None else []
    finally:
        wb.close()
    p = _busca(filas, lambda t: t.startswith('cuentas por cobrar'), 10)
    if p is None:
        raise ValueError("%s: no «Cuentas por Cobrar» header on the Resumen sheet" % filename)
    ih = next(i for i in range(p[0] + 1, min(p[0] + 4, len(filas))) if any(_norm(v) == 'total' for v in filas[i]))
    cols = [(j, str(v).strip()) for j, v in enumerate(filas[ih]) if isinstance(v, str) and v.strip()]
    jt = next(j for j, n in cols if _norm(n) == 'total')
    presta = [(j, n) for j, n in cols if j != jt]
    jn = min(j for j, _ in presta) - 2         # the borrower names sit two columns left of the first lender
    jn = next((j for j in range(jn, -1, -1) if any(isinstance(_celda(filas, i, j), str) for i in range(ih + 1, len(filas)))), jn)
    filas_m, total, bloque = [], None, []
    for i in range(ih + 1, len(filas)):
        n = _celda(filas, i, jn)
        if not (isinstance(n, str) and n.strip()):
            if bloque:
                filas_m.append(bloque)
                bloque = []
            continue
        if _norm(n).startswith('notas'):
            break
        vals = [_num(_celda(filas, i, j)) or 0.0 for j, _ in presta]
        fila = {'n': n.strip(), 'k': clave_desarrollo(n), 'v': [round(x / 1e6, 6) for x in vals], 't': _mdp(_celda(filas, i, jt)) or 0.0}
        if _norm(n) == 'total':
            if bloque:
                filas_m.append(bloque)
                bloque = []
            total = fila
            break
        bloque.append(fila)
    if bloque:
        filas_m.append(bloque)
    if total is None:
        raise ValueError("%s: no Total row on the Resumen sheet" % filename)
    ok, lines, deudores = True, [], []
    for b in filas_m:
        cab, resto = b[0], b[1:]
        if resto and abs(cab['t'] - sum(x['t'] for x in resto)) < 2e-5:
            deudores.append(dict(cab, hijos=resto))
        else:
            deudores.extend(dict(x, hijos=[]) for x in b)
    for d in deudores:
        s = sum(d['v'])
        if abs(s - d['t']) > 2e-5:
            ok = False
            lines.append('%-30s suma de prestamistas %.4f vs Total %.4f' % (d['n'][:30], s, d['t']))
    for k, (_, n) in enumerate(presta):
        s = sum(d['v'][k] for d in deudores)
        c = abs(s - total['v'][k]) < 1e-4
        ok &= c
        lines.append('Prestamista %-28s suma %.4f vs Total %.4f: dif %.6f mdp' % (n[:28], s, total['v'][k], abs(s - total['v'][k])))
    s = sum(d['t'] for d in deudores)
    ok &= abs(s - total['t']) < 1e-4
    movs, ltot = _ledger(integ)
    if ltot is not None:
        sl = sum(m['v'] for m in movs)
        ok &= abs(sl - ltot) < 1e-4 and abs(ltot - total['t']) < 1e-4
        lines.append('Integración: suma %.4f, Total %.4f, Resumen %.4f: dif %.6f mdp'
                     % (sl, ltot, total['t'], max(abs(sl - ltot), abs(ltot - total['t']))))
    om, otot = _ledger(otros)
    titulo_otros = next((str(r[0]).strip() for r in otros[:2] if isinstance(r[0], str)), None)
    data = {'archivo': filename, 'asOf': as_of, 'prestamistas': [n for _, n in presta], 'deudores': deudores,
            'total': total, 'movimientos': movs, 'otros': {'titulo': titulo_otros, 'movimientos': om, 'total': otot} if om else None}
    return data, bool(ok), lines


# ── «Necesidad de Capital al aa-mm-dd.xlsx» (weekly) ────────────────────────
# The visible sheet (older workbooks keep each past week hidden): on the left,
# the projects whose cash covers this week's payments (Pagos, Saldos, Saldos
# después de pagos) closed by a Total row; on the right, since 2025, the ones
# that need an intercompany loan (Pagos, A cubrir con saldos, Necesidad de
# préstamo). The 2024 layout breaks the right side down by expense line instead;
# only its left table is read.
def parse_necesidad(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        # the week's sheet: «Necesidades» when it is there, else the last visible
        # one (2024 named them by date, «030424») with a «Proyecto / Plaza» header
        visibles = [ws for ws in wb.worksheets if getattr(ws, 'sheet_state', 'visible') == 'visible'] or wb.worksheets
        cands = [ws for ws in visibles if _norm(ws.title) == 'necesidades'] + list(reversed(visibles))
        ws, filas, p = None, [], None
        for c in cands:
            f = _filas(c, max_col=30, max_row=min(c.max_row or 300, 300))
            p = _busca(f, lambda t: t.startswith('proyecto / plaza') or t == 'proyecto/plaza', 12)
            if p is not None:
                ws, filas = c, f
                break
    finally:
        wb.close()
    if p is None:
        raise ValueError("%s: no sheet with the weekly «Proyecto / Plaza» table — not the usual Necesidad de Capital layout" % filename)
    ih, j0 = p
    cab = [_norm(v) if isinstance(v, str) else '' for v in filas[ih]]
    ok, lines = True, []

    def tabla(j0, ncols, jdesc=None):
        out, tot = [], None
        for r in filas[ih + 1:]:
            n = r[j0] if j0 < len(r) else None
            if not (isinstance(n, str) and n.strip()):
                if out or tot:
                    if tot:
                        break
                continue
            vals = [_num(r[j0 + k]) if j0 + k < len(r) else None for k in range(1, ncols + 1)]
            if all(v is None for v in vals) or _norm(n).startswith('subtotal'):
                continue                  # a «Subtotal» inside the table is part of the Total too
            if _norm(n) == 'total':
                tot = vals
                break
            d = r[jdesc] if jdesc is not None and jdesc < len(r) else None
            out.append((n.strip(), vals + [d.strip() if isinstance(d, str) else None]))
        return out, tot

    izq, tot_i = tabla(j0, 3)
    if not izq or tot_i is None:
        raise ValueError("%s: the projects table has no rows or no Total" % filename)
    for k, nom in enumerate(('pagos', 'saldos', 'despues')):
        s = sum(v[k] or 0 for _, v in izq)
        c = _cuadra(s, tot_i[k], 2.0) or tot_i[k] is None      # a blank in the Total row declares nothing
        ok &= c
        lines.append('Proyectos %-8s suma %.2f vs Total %s: dif %.2f' % (nom, s, tot_i[k], abs(s - (tot_i[k] or 0))))
    malas = [n for n, v in izq if v[0] is not None and v[1] is not None and v[2] is not None and abs(v[1] - v[0] - v[2]) > 2.0]
    if malas:      # reported, not held against the file: its Total rows are the check
        lines.append('Aviso: saldos − pagos no da el saldo después de pagos en %s' % ', '.join(malas[:6]))
    # the loan-need table (2025 on): its own «Proyecto / Plaza» header to the right, without a «Rubro» column
    der, tot_d = [], None
    jn = next((j for j in range(j0 + 4, len(cab)) if cab[j].startswith('proyecto')), None)
    if jn is not None and not (jn + 1 < len(cab) and cab[jn + 1] == 'rubro'):
        jdesc = jn + 4 if jn + 4 < len(cab) and cab[jn + 4].startswith('descripci') else None
        der, tot_d = tabla(jn, 3, jdesc)
        if der and tot_d is not None:
            for k, nom in enumerate(('pagos', 'cubrir', 'necesidad')):
                s = sum(v[k] or 0 for _, v in der)
                c = _cuadra(s, tot_d[k], 2.0) or tot_d[k] is None
                ok &= c
                lines.append('Necesidad %-9s suma %.2f vs Total %s: dif %.2f' % (nom, s, tot_d[k], abs(s - (tot_d[k] or 0))))
    m = lambda v: round(v / 1e6, 4) if v is not None else None
    data = {'archivo': filename, 'asOf': as_of, 'hoja': ws.title,
            'cubiertos': [{'n': n, 'k': clave_desarrollo(n), 'pagos': m(v[0]), 'saldos': m(v[1]), 'despues': m(v[2])} for n, v in izq],
            'total': {'pagos': m(tot_i[0]), 'saldos': m(tot_i[1]), 'despues': m(tot_i[2])},
            'necesitan': [{'n': n, 'k': clave_desarrollo(n), 'pagos': m(v[0]), 'cubrir': m(v[1]), 'necesidad': m(v[2]),
                           'nota': v[3]} for n, v in der],
            'total_necesidad': {'pagos': m(tot_d[0]), 'cubrir': m(tot_d[1]), 'necesidad': m(tot_d[2])} if tot_d else None}
    return data, bool(ok), lines


# ── «Reinversión de utilidades.xlsx» ────────────────────────────────────────
# Three versions of when each coto's profits can be taken out for reinvestment,
# month by month to Dec-2035, in thousands of pesos. «Utilidades V1» and «V2»
# release a coto's cash once it passes 80% of its sales signed; «V3» counts
# every positive balance. Each version sheet is a run of blocks (a title row,
# then «Proyecto» and the months, then one row per coto or project and a Total).
# «Validaciones» compares the versions: the indicators, and the month-by-month
# equity to recover. Its final amount per version is the check.
def _bloques_version(filas):
    """{title: {'meses': [...], 'filas': [(name, [values])]}} for every block of a version sheet."""
    out, i = {}, 0
    while i < len(filas) - 1:
        t = filas[i][0]
        sig = filas[i + 1]
        if isinstance(t, str) and t.strip() and isinstance(sig[0], str) and _norm(sig[0]) == 'proyecto':
            meses = [(j, v.strftime('%Y-%m')) for j, v in enumerate(sig) if isinstance(v, dt.datetime)]
            rows, k, cerrado = [], i + 2, False
            while k < len(filas) and isinstance(filas[k][0], str) and filas[k][0].strip():
                if not cerrado:          # rows after the Total («Excedente 10 MDP») are notes
                    rows.append((filas[k][0].strip(), [_num(filas[k][j]) or 0.0 for j, _ in meses]))
                    cerrado = _norm(filas[k][0]).startswith('total')
                k += 1
            out[_norm(t)] = {'titulo': t.strip(), 'meses': [m for _, m in meses], 'filas': rows}
            i = k
        else:
            i += 1
    return out


def parse_reinversion(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        va = _hoja(wb, 'validaciones')
        if va is None:
            raise ValueError("%s has no 'Validaciones' sheet" % filename)
        val = _filas(va, max_col=130)
        versiones = {}
        for ws in wb.worksheets:
            m = re.search(r'v(\d)\s*$', _norm(ws.title))
            if m and 'utilidades' in _norm(ws.title):
                versiones['V' + m.group(1)] = _bloques_version(_filas(ws, max_col=130))
    finally:
        wb.close()
    if not versiones:
        raise ValueError("%s: no «Utilidades V1/V2/V3» sheets" % filename)
    # indicators: «Indicador | V1 | V2 | V3»
    pi = _busca(val, lambda t: t == 'indicador', 10)
    indicadores = []
    if pi:
        cols = [(j, str(v).strip()) for j, v in enumerate(val[pi[0]]) if isinstance(v, str) and re.match(r'^v\d$', _norm(v))]
        for r in val[pi[0] + 1:]:
            if not (isinstance(r[0], str) and r[0].strip()) or all(r[j] is None for j, _ in cols):
                break
            indicadores.append({'n': r[0].strip(), 'v': {n: _mdp(r[j], 1e3) for j, n in cols}})
    # month by month equity to recover, per version («Comparativa de Resultados»)
    ps = _busca(val, lambda t: t == 'escenario')
    serie = None
    if ps:
        meses = [(j, v.strftime('%Y-%m')) for j, v in enumerate(val[ps[0]]) if isinstance(v, dt.datetime)]
        vs = {}
        for r in val[ps[0] + 1:ps[0] + 6]:
            m = re.search(r'(\d)$', _norm(r[0])) if isinstance(r[0], str) else None
            if m:
                vs['V' + m.group(1)] = [round((_num(r[j]) or 0.0) / 1e3, 3) for j, _ in meses]
        serie = {'meses': [m for _, m in meses], 'v': vs}
    # per coto and per project, from each version sheet: the block of equity
    # released (V1/V2: «Equity por recuperar acumulado por coto»; V3: «Importe disponible por coto»)
    ok, lines = True, []

    def bloque(bl, *prefijos):
        return next((b for t, b in bl.items() for p in prefijos if t.startswith(p)), None)

    cotos, proyectos = {}, {}
    for v, bl in sorted(versiones.items()):
        bc = bloque(bl, 'equity por recuperar acumulado por coto', 'importe disponible por coto')
        bp = bloque(bl, 'equity por recuperar acumulado por proyecto', 'importe disponible por proyecto')
        bs = bloque(bl, 'saldo final de efectivo')
        if bc is None:
            continue
        fin = {}
        for n, xs in (bs['filas'] if bs else []):
            fin[n] = xs[-1] if xs else 0.0
        for n, xs in bc['filas']:
            if _norm(n).startswith('total'):
                continue
            c = cotos.setdefault(n, {'n': n, 'k': clave_desarrollo(n), 'v': {}})
            primero = next((bc['meses'][j] for j, x in enumerate(xs) if abs(x) > 0.05), None)
            c['v'][v] = {'libera': round(xs[-1] / 1e3, 3) if xs else 0.0, 'desde': primero,
                         'saldo': round(fin.get(n, 0.0) / 1e3, 3)}
        if bp:
            tot = next((xs for n, xs in bp['filas'] if _norm(n).startswith('total')), None)
            cuerpo = [(n, xs) for n, xs in bp['filas'] if not _norm(n).startswith('total')]
            for n, xs in cuerpo:
                proyectos.setdefault(n, {'n': n, 'k': clave_desarrollo(n), 'v': {}})['v'][v] = [round(x / 1e3, 3) for x in xs]
            if tot:
                s = sum(xs[-1] for _, xs in cuerpo)
                c = abs(s - tot[-1]) < 2.0
                ok &= c
                lines.append('%s por proyecto: suma a dic %.1f vs Total %.1f: dif %.1f' % (v, s / 1e3, tot[-1] / 1e3, abs(s - tot[-1]) / 1e3))
                if serie and v in serie['v'] and serie['v'][v]:
                    d = abs(serie['v'][v][-1] - tot[-1] / 1e3)
                    ok &= d < 0.01
                    lines.append('%s Validaciones vs hoja de la versión, a dic: dif %.3f mdp' % (v, d))
            if 'meses' not in proyectos.get('_', {}):
                proyectos['_'] = {'meses': bp['meses']}
    meses_p = proyectos.pop('_', {}).get('meses')
    if not cotos:
        raise ValueError("%s: no per-coto block in the version sheets" % filename)
    notas = [str(r[0]).strip() for r in val if isinstance(r[0], str) and re.match(r'^\d+\.\s', r[0].strip())]
    data = {'archivo': filename, 'asOf': as_of, 'versiones': sorted(versiones), 'indicadores': indicadores,
            'serie': serie, 'cotos': list(cotos.values()), 'proyectos': {'meses': meses_p, 'filas': list(proyectos.values())},
            'notas': notas}
    return data, bool(ok), lines


# ── «TABLA RESERVAS TERRITORIALES AT.xlsx» ──────────────────────────────────
# One record per site, several rows tall: the name in column B starts it, the
# rows below carry on its texts (address and map link, density, market by
# product, comments). Status «1 - Definir Producto» … «6 - No Procede» (the
# «Variables» sheet). There are no totals to check: a status that is filled in
# must be one of the six (a site just added may have none yet).
_ESTATUS = {1: 'Definir producto', 2: 'Proforma de negocios', 3: 'Carta de intención', 4: 'Convenio de negocios',
            5: 'Firma de fideicomiso', 6: 'No procede'}
_CAMPOS = {'proyecto': 'n', 'estatus': 'estatus', 'plaza': 'plaza', 'responsable': 'responsable', 'ubicacion': 'ubicacion',
           'densidad autorizada': 'densidad', 'superficie total': 'superficie', 'valor total': 'valor', 'valor m2': 'valor_m2',
           'aportacion (%)': 'aportacion', 'garantia': 'garantia', 'mercado': 'mercado', 'comentarios': 'comentarios',
           'estrategias a seguir': 'estrategias', 'comentarios ranman': 'comentarios_ranman'}


def _m2(texto):
    """'29,967.02 m2' -> 29967.02; '20.51 has' -> 205100; None when it gives no surface."""
    t = _norm(texto or '').replace(',', '')
    m = re.search(r'(\d+(?:\.\d+)?)\s*(m2|m²|mts|has?\b|hect)', t)
    if not m:
        return None
    v = float(m.group(1))
    return round(v * 10000 if m.group(2).startswith('h') else v, 2)


def _mdp_texto(texto):
    """The first «$ 39 MDP» / «$ 107.2 MDP» amount in a text, in mdp."""
    m = re.search(r'\$?\s*(\d+(?:[.,]\d+)?)\s*m[dp]p', _norm(texto or '').replace(',', ''))
    return float(m.group(1)) if m else None


def parse_reservas(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = _hoja(wb, 'reservas territoriales') or wb.worksheets[0]
        filas = _filas(ws, max_col=20)
    finally:
        wb.close()
    p = _busca(filas, lambda t: t == 'estatus', 10)
    if p is None:
        raise ValueError("%s: no ESTATUS header" % filename)
    ih = p[0]
    cols = {}
    for j, v in enumerate(filas[ih]):
        k = _CAMPOS.get(_norm(v)) if isinstance(v, str) else None
        if k:
            cols[k] = j
    if 'n' not in cols:
        raise ValueError("%s: no PROYECTO header" % filename)
    sitios, actual, grupo = [], None, None
    for r in filas[ih + 1:]:
        a = r[0] if r else None
        if isinstance(a, str) and a.strip():
            grupo = a.strip()
        n = r[cols['n']] if cols['n'] < len(r) else None
        if isinstance(n, str) and n.strip():
            actual = {'grupo': grupo, **{k: [] for k in cols}}
            sitios.append(actual)
        if actual is None:
            continue
        for k, j in cols.items():
            v = r[j] if j < len(r) else None
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            actual[k].append(v.strip() if isinstance(v, str) else v)
    ok, lines, out = True, [], []
    for s in sitios:
        est = ' '.join(str(x) for x in s.get('estatus', []))
        m = re.match(r'\s*(\d)', est)
        num = int(m.group(1)) if m and int(m.group(1)) in _ESTATUS else None
        if num is None and est.strip():
            ok = False
            lines.append('%s: estatus no reconocido (%r) — dif 1' % (' '.join(map(str, s['n']))[:40], est[:40]))
        elif num is None:
            lines.append('%s: sin estatus en el archivo' % ' '.join(map(str, s['n']))[:40])
        texto = lambda k: '\n'.join(str(x) for x in s.get(k, []))
        ubic = [str(x) for x in s.get('ubicacion', [])]
        mapa = next((u for u in ubic if u.startswith('http')), None)
        apor = s.get('aportacion', [])
        ap = next((x for x in apor if isinstance(x, (int, float))), None)
        if ap is None:
            mp = next((re.search(r'(\d+(?:\.\d+)?)\s*%', str(x)) for x in apor if re.search(r'\d\s*%', str(x))), None)
            ap = float(mp.group(1)) / 100 if mp else None
        valor = texto('valor')
        inicial = next((x for x in s.get('valor', []) if 'inicial' in _norm(str(x))), None)
        ranman = next((x for x in s.get('valor', []) if 'ranman' in _norm(str(x)) or 'aceptada' in _norm(str(x))), None)
        plaza = next((str(x) for x in s.get('plaza', []) if str(x).strip()), None) or s['grupo'] or ''
        out.append({'n': ' '.join(str(x) for x in s['n']).strip(), 'grupo': s['grupo'], 'plaza': plaza.strip(),
                    'estatus': num, 'estatus_txt': est.strip(), 'responsable': texto('responsable'),
                    'ubicacion': '\n'.join(u for u in ubic if not u.startswith('http')), 'mapa': mapa,
                    'densidad': texto('densidad'), 'superficie_txt': texto('superficie'), 'm2': _m2(texto('superficie')),
                    'valor_txt': valor, 'valor_inicial': _mdp_texto(inicial), 'valor_ranman': _mdp_texto(ranman),
                    'valor_m2': next((x for x in s.get('valor_m2', []) if isinstance(x, (int, float))), None),
                    'aportacion': ap if ap is None or ap <= 1 else ap / 100,
                    'garantia': texto('garantia'), 'mercado': texto('mercado'), 'comentarios': texto('comentarios'),
                    'estrategias': texto('estrategias'), 'comentarios_ranman': texto('comentarios_ranman')})
    if not out:
        raise ValueError("%s: no sites found under the header" % filename)
    lines.append('%d sitios, %d con estatus reconocido' % (len(out), sum(1 for x in out if x['estatus'])))
    data = {'archivo': filename, 'asOf': as_of, 'estatus': [[k, v] for k, v in _ESTATUS.items()], 'sitios': out}
    return data, bool(ok), lines
