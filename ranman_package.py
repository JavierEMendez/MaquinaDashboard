"""RANMAN's monthly package -> what each file is, and the parsers for the ones
the dashboard reads beyond Flujo y PLP and the Cuadro de Riesgos.

RANMAN uploads its files to OneDrive, under
    Archivos Ranman/<año>/<n>_<Mes>/<n> <entregable>/<archivo>
The folder numbers have drifted over the years (Breakeven was 10 in 2024 and is
8 now; the signed statements have been 9 through 15), so a file is recognised
by its NAME, never by its folder. `clasifica()` does that for the app, the
upload page and tools/ranman_sync.py alike.

Parsers in this module (each returns (data, ok, lines) — ok False means the
file's own totals don't add up and nothing should be stored):
    parse_edo_resultados  «Estado de Resultados al dd-mm-aa.xlsx»   lifetime P&L by etapa
    parse_breakeven       «Breakeven al dd-mm-aa.xlsx»              units / pesos to reach UAIR 0% and 10%
    parse_pipeline        «Pipeline al dd-mm-aa.xlsx»               units still to sell, by coto and quarter
    parse_sabana          «NNN PLAN OPERATIVO <año>_<MES>_….xlsx»   plan vs real by development, month by month
Amounts come in pesos (Edo. Resultados) or thousands (Breakeven, Pipeline
prices) and are stored in mdp, like the other Ranman tableros.
"""
import calendar
import datetime as dt
import io
import re
import unicodedata

import openpyxl

import ranman_deuda_parser


def _norm(s) -> str:
    """Lower case, no accents, single spaces — how names are compared."""
    s = unicodedata.normalize('NFKD', str(s or ''))
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', s).strip().lower()


# ── developments: one key per desarrollo, shared by every report ────────────
PLAZAS = {'SLP': 'San Luis Potosí', 'QRO': 'Querétaro', 'AGS': 'Aguascalientes', 'CEL': 'Celaya',
          'SMA': 'San Miguel de Allende', 'GDL': 'Guadalajara', 'CAN': 'Cancún'}
DESARROLLOS = {
    'puerta_natura': ('Puerta Natura', 'SLP'), 'cantabria': ('Cantabria', 'SLP'),
    'kima': ('Kima (Gran Peñón)', 'SLP'), 'trivee': ('Trivé (La Pila)', 'SLP'),
    'valpongero': ('Valpongero', 'SLP'), 'satelite': ('Cd. Satélite', 'SLP'),
    'azhala': ('Azhala / Alella', 'QRO'), 'cuspide': ('Cúspide', 'QRO'), 'mompani': ('Mompani', 'QRO'),
    'puerta_piedra_qro': ('Puerta de Piedra QRO', 'QRO'),
    'grand_natura': ('Grand Natura', 'AGS'), 'cartuja': ('La Cartuja', 'AGS'), 'vento': ('Vento', 'AGS'),
    'carmel': ('Carmel', 'AGS'), 'vinedos_elizondo': ('Viñedos Elizondo', 'AGS'),
    'country_golf': ('Country Golf', 'AGS'),
    'miranda': ('Miranda', 'CEL'), 'cartagena': ('Alcázar / Cartagena', 'CEL'),
    'veredas': ('Las Veredas', 'SMA'), 'puerta_piedra': ('Puerta de Piedra', 'GDL'),
    'punta_norte': ('Punta Norte', 'CAN'),
}
# First match wins, so the more specific phrases go first.
_CLAVES = [('vinedos elizondo', 'vinedos_elizondo'), ('puerta natura', 'puerta_natura'),
           ('puerta de piedra qro', 'puerta_piedra_qro'), ('qro - puerta de piedra', 'puerta_piedra_qro'),
           ('puerta de piedra', 'puerta_piedra'), ('cantabria', 'cantabria'), ('ximonco', 'cantabria'),
           ('kima', 'kima'), ('gran penon', 'kima'), ('trivee', 'trivee'), ('trive', 'trivee'),
           ('la pila', 'trivee'), ('valpongero', 'valpongero'), ('satelite', 'satelite'),
           ('azhala', 'azhala'), ('alella', 'azhala'), ('zibata', 'azhala'), ('cuspide', 'cuspide'),
           ('mompani', 'mompani'), ('grand natura', 'grand_natura'), ('gran natura', 'grand_natura'),
           ('cartuja', 'cartuja'), ('vento', 'vento'), ('carmel', 'carmel'), ('country golf', 'country_golf'),
           ('miranda', 'miranda'), ('cartagena', 'cartagena'), ('alcazar', 'cartagena'),
           ('veredas', 'veredas'), ('punta norte', 'punta_norte'), ('cancun', 'punta_norte')]


def clave_desarrollo(texto):
    """'QRO - AZHALA / ALELLA', 'Azhala C02', 'Alella' -> 'azhala'; None when it names no development."""
    t = _norm(texto)
    for frase, clave in _CLAVES:
        if frase in t:
            return clave
    return None


# ── the 15 deliverables and how their files are named ───────────────────────
# (key, folder number today, name, regex on the normalised file name, how the
# date is written, parsed by the dashboard). The regexes are written so that
# JavaScript accepts them unchanged — the upload page filters a whole month
# folder with them before sending anything.
ENTREGABLES = [
    ('flujo',          1,  'Resumen de Flujos-PLP',        r'^flujo y plp dra - maquina - \d{2}-\d{2}-\d{2}\.xlsx$', 'dma', True),
    ('edo_resultados', 2,  'Edo. Resultados por Proyecto', r'^estado de resultados al \d{2}-\d{2}-\d{2}.*\.xlsx$',   'dma', True),
    ('necesidad',      3,  'Necesidad de Capital',         r'^necesidad(es)? de capital al \d{2}-\d{2}-\d{2}.*\.xlsx$', 'amd', False),
    ('riesgos',        4,  'Cuadro de Riesgos',            r'^cuadro de riesgos ranman .*\.xls[xm]$',                'riesgos', True),
    ('bp',             5,  'BPs',                          r'^(\d{2} )?bp .*\.xlsx$',                                None, False),
    ('pipeline',       6,  'Pipeline',                     r'^pipeline al \d{2}-\d{2}-\d{2}.*\.xlsx$',               'dma', True),
    ('aaa',            7,  'AAA',                          r'^reporte aaa .*\.pdf$',                                 'mes', True),
    ('breakeven',      8,  'Breakeven',                    r'^breakeven al \d{2}-\d{2}-\d{2}.*\.xlsx$',              'dma', True),
    ('reservas',       9,  'Reservas Territoriales',       r'^tabla reservas territoriales.*\.xlsx$',                None, False),
    ('sabana',         10, 'Sábana Operativa',             r'plan operativo \d{4}_[a-z]+_.*\.xlsx$',                 'mes', True),
    ('aportaciones',   11, 'Aportaciones MAQUINA',         r'^aportaciones maquina al \d{2}-\d{2}-\d{2}.*\.xlsx$',   'dma', False),
    ('prestamos',      12, 'Préstamo entre proyectos',     r'^control de prestamos entre proyectos.*\.xlsx$',        'dma', False),
    ('ccc',            13, 'Control de Créditos CCC',      r'^control de ccc al \d{2}-\d{2}-\d{2}.*\.xlsx$',         'dma', False),
    ('reinversion',    14, 'Reinversión de Utilidades',    r'^reinversion de utilidades.*\.xlsx$',                   None, False),
    ('eeff',           15, 'Estados Financieros Firmados', r'(estados financieros|eeff).*\.pdf$',                    'mes', False),
    ('analiticas',     15, 'Analíticas (balanza)',         r'^(\d+\. )?(dra )?analiticas.*\.pdf$',                   'mes', False),
]
ENTREGABLE = {e[0]: e for e in ENTREGABLES}
PARSEADOS = {e[0] for e in ENTREGABLES if e[5]}
# What tools/ranman_sync.py and the upload page send. BPs wait for their parser:
# the models run to 14 MB and the backup trees hold some 240 more.
ENVIABLES = {e[0] for e in ENTREGABLES if e[0] != 'bp'}

OMITE_NOMBRE = ('~$', '(previo)', 'previo ---', ' - copia', 'respaldo')
# Sub-folders that hold backups, scenarios and supporting paperwork, never a
# deliverable (the Cancún BPs spell theirs «Respalos»).
OMITE_CARPETA = ('respaldo', 'respalos', 'tarjetas', 'presentacion', 'nuevos proyectos', 'versiones',
                  'escenarios', 'alterno', 'opciones', 'claude outputs', '_tablero', 'soporte', 'complementos')

MESES_ES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6, 'julio': 7,
            'agosto': 8, 'septiembre': 9, 'setiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12}
MESES_EN = {'january': 1, 'february': 2, 'march': 3, 'april': 4, 'may': 5, 'june': 6, 'july': 7,
            'august': 8, 'september': 9, 'october': 10, 'november': 11, 'december': 12}


def _fin_de_mes(anio: int, mes: int) -> str:
    return dt.date(anio, mes, calendar.monthrange(anio, mes)[1]).isoformat()


def _fecha(modo, nombre_norm: str, original: str):
    """The cut date a file name carries, as ISO, or None."""
    try:
        if modo in ('dma', 'amd'):
            m = re.search(r'(\d{2})-(\d{2})-(\d{2})', nombre_norm)
            if not m:
                return None
            a, b, c = (int(x) for x in m.groups())
            d, mes, anio = (a, b, c) if modo == 'dma' else (c, b, a)
            return dt.date(2000 + anio, mes, d).isoformat()
        if modo == 'riesgos':
            return ranman_deuda_parser.fecha_de_nombre(original)
        if modo == 'mes':
            # «Reporte AAA Agosto 26», «EEFF DRA al 31 de Marzo de 2026», «PLAN OPERATIVO 2026_JULIO_...»
            m = re.search(r'(\d{2})-(\d{2})-(\d{2})', nombre_norm)
            if m:
                d, mes, anio = (int(x) for x in m.groups())
                return dt.date(2000 + anio, mes, d).isoformat()
            # the Sábana names its month right after the year («2026_JULIO_12 agosto 2026_…»)
            m = re.search(r'(20\d{2})_([a-z]+)', nombre_norm)
            if m and m.group(2) in MESES_ES:
                return _fin_de_mes(int(m.group(1)), MESES_ES[m.group(2)])
            mes = next((v for k, v in MESES_ES.items() if re.search(r'(^|[^a-z])%s([^a-z]|$)' % k, nombre_norm)), None)
            anio = re.search(r'(?<!\d)(20\d{2})(?!\d)', nombre_norm) or re.search(r'(?<![\d-])(\d{2})(?![\d-])', nombre_norm)
            if mes and anio:
                y = int(anio.group(1))
                return _fin_de_mes(y if y > 2000 else 2000 + y, mes)
    except (ValueError, TypeError):
        return None
    return None


def omitir_ruta(ruta: str) -> bool:
    """True for files inside backup / scenario / paperwork sub-folders."""
    partes = [_norm(p) for p in re.split(r'[\\/]+', ruta or '')[:-1]]
    return any(o in p for p in partes for o in OMITE_CARPETA)


def clasifica(nombre: str):
    """File name -> (kind, as_of ISO or None); (None, None) when it isn't a deliverable."""
    base = re.split(r'[\\/]', nombre or '')[-1]
    n = _norm(base)
    if not n or any(o in n for o in OMITE_NOMBRE):
        return None, None
    for clave, _num, _et, rx, modo, _p in ENTREGABLES:
        if re.search(rx, n):
            return clave, (_fecha(modo, n, base) if modo else None)
    return None, None


def periodo_de_ruta(ruta: str):
    """'2026/8_August/4 Cuadro de Riesgos/x.xlsm' -> '2026-08' (the reporting month of the folder)."""
    partes = [p for p in re.split(r'[\\/]+', ruta or '') if p]
    for i in range(len(partes) - 1):
        if re.fullmatch(r'20\d{2}', partes[i]):
            m = re.match(r'(\d{1,2})_', partes[i + 1])
            if m and 1 <= int(m.group(1)) <= 12:
                return '%s-%02d' % (partes[i], int(m.group(1)))
    return None


def patrones_js():
    """[[kind, label, regex]] for the upload page's folder filter."""
    return [[k, et, rx] for k, _n, et, rx, _m, _p in ENTREGABLES]


# ── helpers shared by the parsers ───────────────────────────────────────────
def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _hoja(wb, *candidatos):
    """The sheet whose normalised name equals (or starts with) one of the candidates."""
    nombres = {_norm(s): s for s in wb.sheetnames}
    for c in candidatos:
        if c in nombres:
            return wb[nombres[c]]
    for c in candidatos:
        for k, s in nombres.items():
            if k.startswith(c):
                return wb[s]
    return None


# ── «Estado de Resultados al dd-mm-aa.xlsx» ─────────────────────────────────
# Lifetime P&L of every etapa (real a la fecha + por ejercer), grouped by plaza.
# Sheet «Edo Resultados»: row 5 is the header; columns C:S give each line as %
# of sales and V:AL the same lines in pesos. Rows 1-2 hold Ranman's own green
# and red limits per line. Each plaza is a run of etapa rows closed by its
# subtotal; the grand total sits alone after the last plaza.
CONCEPTOS = [('ventas', 'Ventas'), ('terreno', 'Terreno'), ('permisos', 'Permisos y Licencias'),
             ('estudios', 'Estudios y Proyectos'), ('urbanizacion', 'Urbanización'),
             ('infraestructura', 'Infraestructura'), ('edificacion', 'Edificación'),
             ('margen_bruto', 'Margen Bruto'), ('admin_plaza', 'Admin (Plaza)'),
             ('admin_corp', 'Admin (Corporativo)'), ('comercial', 'Comercial'), ('comisiones', 'Comisiones'),
             ('ebit', 'EBIT'), ('fin_operativos', 'Financieros Operativos'), ('fin_lp', 'Financieros LP'),
             ('uair', 'UAIR')]


def parse_edo_resultados(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = _hoja(wb, 'edo resultados')
        if ws is None:
            raise ValueError("%s has no 'Edo Resultados' sheet" % filename)
        filas = list(ws.iter_rows(min_row=1, max_row=120, max_col=60, values_only=True))
    finally:
        wb.close()
    hr = next((i for i, r in enumerate(filas) if len(r) > 1 and _norm(r[1]) == 'estado de resultados'), None)
    if hr is None:
        raise ValueError("%s: no header row with 'Estado de Resultados' in column B" % filename)
    cab = [_norm(v) for v in filas[hr]]
    bloques = [j for j, v in enumerate(cab) if v == 'estado de resultados']
    if len(bloques) < 2:
        raise ValueError("%s: expected the %%-of-sales block and the pesos block side by side" % filename)
    izq, der = bloques[0], bloques[1]
    col_avance = next((j for j in range(izq, der) if cab[j] == 'avance en firmas'), None)

    def col(concepto, desde, hasta):
        t = _norm(concepto)
        return next((j for j in range(desde, hasta) if cab[j] == t), None)
    col_pct = {k: col(et, izq, der) for k, et in CONCEPTOS}
    col_abs = {k: col(et, der, len(cab)) for k, et in CONCEPTOS}
    falta = [et for k, et in CONCEPTOS if col_abs[k] is None]
    if falta:
        raise ValueError("%s: the pesos block lacks %s" % (filename, ", ".join(falta)))

    umbrales = {}
    for r in filas[:hr]:
        t = _norm(r[1]) if len(r) > 1 else ''
        if t in ('verde', 'rojo'):
            for k, _ in CONCEPTOS:
                j = col_pct[k]
                v = _num(r[j]) if j is not None else None
                if v is not None and k != 'ventas':
                    umbrales.setdefault(k, {})[t] = round(v, 4)
    for k, u in umbrales.items():      # margins: more is better; costs: less is better
        if 'verde' in u and 'rojo' in u:
            u['mejor'] = 'alto' if u['verde'] > u['rojo'] else 'bajo'

    def valores(r):
        return {k: round((_num(r[j]) or 0.0) / 1e6, 3) for k, j in col_abs.items()}

    plazas, grupo, del_archivo = [], [], None
    for r in filas[hr + 1:]:
        nombre = r[1] if len(r) > 1 else None
        if nombre is None or not str(nombre).strip():
            if grupo:
                plazas.append(grupo)
                grupo = []
            if del_archivo is None and _num(r[col_abs['ventas']]) is not None and len(plazas) > 0:
                del_archivo = valores(r)       # the file only fills sales, margin, EBIT and UAIR here
            continue
        grupo.append(r)
    if grupo:
        plazas.append(grupo)
    if del_archivo is None:
        raise ValueError("%s: no grand-total row after the last plaza" % filename)

    salida, ok, lines = [], True, []
    for g in plazas:
        sub, etapas = g[-1], g[:-1]
        tot = valores(sub)
        items = []
        for r in etapas:
            v = valores(r)
            av = _num(r[col_avance]) if col_avance is not None else None
            items.append({'n': str(r[1]).strip(), 'k': clave_desarrollo(r[1]),
                          'avance': round(av, 4) if av is not None else None,
                          'pre': (str(r[0] or '').strip() == '*') or (v['ventas'] == 0), 'v': v})
        for k in ('ventas', 'uair'):
            dif = abs(sum(e['v'][k] for e in items) - tot[k]) if items else 0.0
            ok &= dif < 0.05
            lines.append('%-22s %-6s etapas vs subtotal: dif %.3f mdp' % (str(sub[1]).strip()[:22], k, dif))
        salida.append({'n': str(sub[1]).strip(), 'v': tot, 'etapas': items,
                       'avance': round(_num(sub[col_avance]) or 0, 4) if col_avance is not None else None})
    total = {k: round(sum(p['v'][k] for p in salida), 3) for k, _ in CONCEPTOS}
    for k in ('ventas', 'margen_bruto', 'ebit', 'uair'):
        dif = abs(total[k] - del_archivo[k])
        ok &= dif < 0.05
        lines.append('TOTAL %-12s plazas vs total del archivo: dif %.3f mdp' % (k, dif))
    data = {'archivo': filename, 'asOf': as_of, 'conceptos': [[k, et] for k, et in CONCEPTOS],
            'umbrales': umbrales, 'plazas': salida, 'total': total}
    return data, bool(ok), lines


# ── «Breakeven al dd-mm-aa.xlsx» ────────────────────────────────────────────
# Two sheets, «Breakeven UAIR 0%» and «Breakeven UAIR 10%»: units (A:N) and
# thousands of pesos (P:AC) needed each month of the year. Rows come in runs
# separated by blanks: the first row of a run is the plaza, bold rows inside
# it are projects, plain rows are the project's etapas. «Total» closes the
# table; the «Datos BP» inputs below it are not read.
_MESES_CORTO = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
                'septiembre', 'octubre', 'noviembre', 'diciembre']


def _lee_breakeven(ws, filename):
    filas = [(r, [c.value for c in row], bool(row[0].font and row[0].font.b))
             for r, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 160), max_col=40), 1)]
    titulo = ' '.join(str(v) for v in (filas[0][1] if filas else []) if v)
    anio = re.search(r'(20\d{2})', titulo)
    hi = next((i for i, (_, v, _) in enumerate(filas) if _norm(v[0]) == 'proyecto'), None)
    if hi is None:
        raise ValueError("%s: no 'Proyecto' header in sheet %s" % (filename, ws.title))
    cab = [_norm(x) for x in filas[hi][1]]
    otro = next((j for j in range(1, len(cab)) if cab[j] == 'proyecto'), None)   # pesos block

    def meses(desde, hasta):
        idx = [j for j in range(desde, hasta) if cab[j] in _MESES_CORTO]
        tot = next((j for j in range(desde, hasta) if cab[j] == 'total'), None)
        return idx, tot
    mu, tu = meses(1, otro or len(cab))
    mp, tp = meses((otro or 0) + 1, len(cab)) if otro else ([], None)
    if len(mu) != 12:
        raise ValueError("%s: sheet %s should have the 12 months after 'Proyecto'" % (filename, ws.title))

    out, total, nivel_previo = [], None, None
    inicio_grupo = True
    for _, v, negrita in filas[hi + 1:]:
        nombre = v[0]
        if nombre is None or not str(nombre).strip():
            inicio_grupo = True
            continue
        nn = _norm(nombre)
        u = [round(_num(v[j]) or 0.0, 2) for j in mu]
        p = [round((_num(v[j]) or 0.0) / 1000.0, 3) for j in mp] if mp else []
        fila = {'n': str(nombre).strip(), 'u': u, 'ut': round(_num(v[tu]) or sum(u), 2) if tu is not None else round(sum(u), 2),
                'p': p, 'pt': round((_num(v[tp]) or 0.0) / 1000.0, 3) if tp is not None else round(sum(p), 3)}
        if nn == 'total':
            total = fila
            break
        fila['nivel'] = 'plaza' if inicio_grupo else ('proyecto' if negrita else 'etapa')
        fila['k'] = clave_desarrollo(nombre) if fila['nivel'] != 'plaza' else None
        out.append(fila)
        inicio_grupo = False
    if total is None:
        raise ValueError("%s: sheet %s has no 'Total' row" % (filename, ws.title))
    return (int(anio.group(1)) if anio else None), out, total


def _valida_breakeven(nombre, filas, total, lines):
    ok = True

    def cmp(etq, hijos, padre):
        nonlocal ok
        if not hijos:
            return
        du = abs(sum(h['ut'] for h in hijos) - padre['ut'])
        dp = abs(sum(h['pt'] for h in hijos) - padre['pt'])
        ok &= du < 0.05 and dp < 0.05
        lines.append('%-4s %-26s unidades dif %.3f · mdp dif %.3f' % (nombre, etq[:26], du, dp))
    plazas = [f for f in filas if f['nivel'] == 'plaza']
    cmp('plazas vs Total', plazas, total)
    i = 0
    while i < len(filas):
        f = filas[i]
        j = i + 1
        hijos = []
        if f['nivel'] == 'plaza':
            while j < len(filas) and filas[j]['nivel'] != 'plaza':
                if filas[j]['nivel'] == 'proyecto':
                    hijos.append(filas[j])
                j += 1
            cmp(f['n'] + ' (proyectos)', hijos, f)
        elif f['nivel'] == 'proyecto':
            while j < len(filas) and filas[j]['nivel'] == 'etapa':
                hijos.append(filas[j])
                j += 1
            cmp(f['n'] + ' (etapas)', hijos, f)
        i += 1
    return ok


def parse_breakeven(file_bytes: bytes, filename: str, as_of: str):
    # fonts are needed to tell projects from etapas, so not read-only (the file is small)
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    try:
        hojas = {'be0': _hoja(wb, 'breakeven uair 0%', 'breakeven uair 0'),
                 'be10': _hoja(wb, 'breakeven uair 10%', 'breakeven uair 10')}
        if not all(hojas.values()):
            raise ValueError("%s needs both 'Breakeven UAIR 0%%' and 'Breakeven UAIR 10%%' sheets" % filename)
        data, ok, lines, anio = {'archivo': filename, 'asOf': as_of}, True, [], None
        for k, ws in hojas.items():
            a, filas, total = _lee_breakeven(ws, filename)
            anio = anio or a
            data[k] = {'filas': filas, 'total': total}
            ok &= _valida_breakeven(k, filas, total, lines)
    finally:
        wb.close()
    data['anio'] = anio
    return data, bool(ok), lines


# ── «Pipeline al dd-mm-aa.xlsx» ─────────────────────────────────────────────
# Sheet «Unidades»: one row per coto — plaza, desarrollo, coto, market, average
# price (thousands), absorption (units a month), project units — then the units
# still to sell at each quarter, 2020 to 2036. Plaza and desarrollo are written
# once per block. A «Total» row closes the table.
def parse_pipeline(file_bytes: bytes, filename: str, as_of: str):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = _hoja(wb, 'unidades')
        if ws is None:
            raise ValueError("%s has no 'Unidades' sheet" % filename)
        filas = list(ws.iter_rows(min_row=1, max_row=260, max_col=140, values_only=True))
    finally:
        wb.close()
    hi = next((i for i, r in enumerate(filas) if 'plaza' in [_norm(v) for v in r[:4]]
               and 'desarrollo' in [_norm(v) for v in r[:4]]), None)
    if hi is None or hi == 0:
        raise ValueError("%s: no PLAZA / DESARROLLO header row — this layout isn't recognised "
                         "(the Jan-2026 file has a different one)" % filename)
    cab = [_norm(v) for v in filas[hi]]

    def c(*nombres):
        return next((j for j, v in enumerate(cab) if v in nombres), None)
    cp, cd, cc, cm = c('plaza'), c('desarrollo'), c('coto / reserva', 'coto'), c('mercado')
    cpr, cab_, cu = c('precio promedio'), c('absorcion promedio'), c('unidades del proyecto')
    anios, anio = filas[hi - 1], None
    trims = []
    for j, v in enumerate(cab):
        y = anios[j] if j < len(anios) else None
        if isinstance(y, (int, float)) and 2000 < y < 2100:
            anio = int(y)
        if re.fullmatch(r'q[1-4]', v or '') and anio:
            trims.append((j, '%d-Q%s' % (anio, v[1])))
    if not trims or None in (cp, cd, cc):
        raise ValueError("%s: the quarterly columns or the PLAZA / DESARROLLO / COTO columns weren't found" % filename)

    out, total, plaza, desa = [], None, None, None
    for r in filas[hi + 1:]:
        if cp is not None and r[cp] not in (None, ''):
            plaza = str(r[cp]).strip()
        if r[cd] not in (None, ''):
            desa = str(r[cd]).strip()
        coto = r[cc]
        vals = [_num(r[j]) for j, _ in trims]
        if _norm(coto) == 'total':
            total = [round(v or 0, 2) for v in vals]
            break
        if coto in (None, '') and not any(v for v in vals):
            continue
        out.append({'plaza': plaza, 'd': desa, 'k': clave_desarrollo(desa), 'coto': str(coto or '').strip(),
                    'mercado': str(r[cm]).strip() if cm is not None and r[cm] else '',
                    'precio': round((_num(r[cpr]) or 0) / 1000.0, 4) if cpr is not None else None,
                    'abs': round(_num(r[cab_]) or 0, 3) if cab_ is not None else None,
                    'u': round(_num(r[cu]) or 0, 2) if cu is not None else None,
                    'r': [round(v or 0, 2) for v in vals]})
    if total is None:
        raise ValueError("%s: no 'Total' row under the cotos" % filename)
    # trim the empty tail: keep up to the first quarter where everything is sold
    ult = max((i for i, v in enumerate(total) if v > 0.005), default=len(total) - 1)
    ult = min(len(total) - 1, ult + 1)
    for f in out:
        f['r'] = f['r'][:ult + 1]
    total = total[:ult + 1]
    lines, ok = [], True
    dif = max(abs(sum(f['r'][i] for f in out) - total[i]) for i in range(len(total)))
    ok &= dif < 0.5
    lines.append('cotos vs Total, por trimestre: dif máx %.2f unidades' % dif)
    data = {'archivo': filename, 'asOf': as_of, 'trimestres': [t for _, t in trims][:ult + 1],
            'filas': out, 'total': total}
    return data, bool(ok), lines


# ── «NNN PLAN OPERATIVO <año>_<MES>_…_FINAL <MES>.xlsx» (Sábana Operativa) ──
# One sheet per plaza («QRO (2026)», « SLP (2026)», «AGS (2026)», «CEL (2026)»).
# Row 1 marks each year's two blocks — PLAN OPERATIVO (the plan) and YTD (real
# for the months flagged H in row 4, Ranman's forecast for those flagged P) —
# row 2 carries the year, row 3 the months. Column B opens a block of rows:
# «RESUMEN <desarrollo>» holds apartados, individualizaciones (or lot-sale
# equivalents), construction starts and completions. The first RESUMEN of a
# plaza with several developments is the plaza's own summary. The INDICADORES
# blocks are left out: their «real» side carries differences, not stocks, and
# reading them needs Ranman to confirm what each one is.
# Firmas here are homes individualised plus lot-sale equivalents; together they
# are the AAA report's «Firmadas» (Aug-2026: 364 + 90 = 454, plan 462).
_SAB_HOJAS = {'qro': 'QRO', 'slp': 'SLP', 'ags': 'AGS', 'cel': 'CEL', 'sma': 'SMA', 'gdl': 'GDL', 'cancun': 'CAN'}
_SAB_UNICO = {'CEL': 'miranda'}          # a plaza whose only RESUMEN is its single development
_SAB_METRICA = [(r'^apartados', 'apartados'), (r'^(individualizaciones|cobranza lotes)', 'firmas'),
                (r'^inicios objetivo', 'inicios'), (r'^terminos objetivo', 'terminos'), (r'^entregas lotes', 'entregas')]


def _sab_num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def parse_sabana(file_bytes: bytes, filename: str, as_of: str):
    anio = int(as_of[:4])
    patron = re.compile(r'^([a-z]+) \(%d\)$' % anio)
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    plazas, desarrollos, lines, ok, meses_reales = [], [], [], True, None
    try:
        hojas = []
        for ws in wb.worksheets:
            m = patron.match(_norm(ws.title))
            if ws.sheet_state == 'visible' and m and m.group(1) in _SAB_HOJAS:
                hojas.append((_SAB_HOJAS[m.group(1)], ws))
        if not hojas:
            raise ValueError("%s: no plaza sheet for %d (expected names like 'QRO (%d)')" % (filename, anio, anio))
        for plaza, ws in hojas:
            r1, r2, _r3, r4 = (list(r) for r in ws.iter_rows(min_row=1, max_row=4, values_only=True))
            pj = next((j for j, v in enumerate(r1) if v == 'PLAN OPERATIVO' and r2[j] == anio), None)
            yj = next((j for j, v in enumerate(r1) if v == 'YTD' and r2[j] == anio), None)
            if pj is None or yj is None:
                raise ValueError("%s: sheet %s has no PLAN OPERATIVO / YTD block for %d" % (filename, ws.title, anio))
            nh = sum(1 for x in r4[yj:yj + 12] if str(x or '').strip().upper() == 'H')
            meses_reales = nh if meses_reales is None else min(meses_reales, nh)
            bloques, actual = [], None
            for row in ws.iter_rows(min_row=5, max_col=yj + 12, values_only=True):
                b = row[1] if len(row) > 1 else None
                c = row[2] if len(row) > 2 else None
                if b not in (None, ''):
                    t = re.sub(r'\s+', ' ', str(b)).strip()
                    actual = {'n': t, 'm': {}} if _norm(t).startswith('resumen') else None
                    if actual:
                        bloques.append(actual)
                if actual is None or c in (None, ''):
                    continue
                et = _norm(c)
                clave = next((k for rx, k in _SAB_METRICA if re.match(rx, et)), None)
                if not clave or clave in actual['m']:
                    continue
                actual['m'][clave] = {'plan': [_sab_num(row[j]) for j in range(pj, pj + 12)],
                                      'real': [_sab_num(row[j]) for j in range(yj, yj + 12)]}
                if clave == 'firmas':
                    actual['lotes'] = 'lotes' in et or 'equiv' in et
            con_datos = [bq for bq in bloques if any(any(v) for m in bq['m'].values() for v in (m['plan'], m['real']))]
            for bq in con_datos:
                nombre = re.sub(r'^resumen\s+', '', bq['n'], flags=re.I).strip()
                k = clave_desarrollo(nombre)
                es_plaza = k is None and len(con_datos) > 1 and plaza not in _SAB_UNICO
                if k is None and plaza in _SAB_UNICO:
                    k = _SAB_UNICO[plaza]
                item = {'plaza': plaza, 'n': nombre, 'k': k, 'lotes': bool(bq.get('lotes')), 'm': bq['m']}
                (plazas if es_plaza else desarrollos).append(item)
    finally:
        wb.close()
    if not desarrollos:
        raise ValueError("%s: no RESUMEN block with %d data" % (filename, anio))
    # The month in the name («…_FINAL JUNIO») is the one that matches the AAA; the
    # H flags in row 4 are sometimes left a month behind (Jun-2026 flags five).
    nh = int(as_of[5:7]) if as_of[:4] == str(anio) else (meses_reales or 0)
    if meses_reales is not None and meses_reales != nh:
        lines.append('row 4 flags %d real months; the file name says %d — using %d' % (meses_reales, nh, nh))

    def ytd(d, met, lado):
        return sum((d['m'].get(met) or {}).get(lado, [0.0] * 12)[:nh])
    # each plaza summary against its developments' homes (lot sales are reported apart)
    for p in plazas:
        hijos = [d for d in desarrollos if d['plaza'] == p['plaza'] and not d['lotes']]
        for met in ('firmas', 'apartados'):
            a, b = ytd(p, met, 'real'), sum(ytd(d, met, 'real') for d in hijos)
            ok &= abs(a - b) <= max(2.0, 0.02 * abs(a))
            lines.append('%-4s %-10s resumen de plaza vs desarrollos, real a la fecha: %g vs %g (dif %g)'
                         % (p['plaza'], met, a, b, a - b))
    total = {met: {lado: sum(ytd(d, met, lado) for d in desarrollos) for lado in ('plan', 'real')}
             for met in ('apartados', 'firmas', 'inicios', 'terminos')}
    data = {'archivo': filename, 'asOf': as_of, 'anio': anio, 'meses_reales': nh,
            'desarrollos': desarrollos, 'plazas': plazas, 'total': total}
    return data, bool(ok), lines
